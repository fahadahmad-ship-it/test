"""
compare.py — structural comparison of the generated workbook against the signed K&K
template. Checks the dimensions the playbook lists: tabs + order, headers, column widths,
row heights, freeze panes, autofilter, gridlines, header/data cell fonts+fills+borders+
alignment, number formats, Summary title font, and the logo. Content (row counts, values,
per-client prose) legitimately differs and is NOT compared.
Run: python3 -I superk/check/compare.py
"""
import os, sys
import openpyxl
OUT = os.path.normpath(os.path.join(os.path.dirname(__file__), "..", "out"))
REF = os.path.normpath(os.path.join(os.path.dirname(__file__), "..", "ref"))
GEN = os.path.join(OUT, "SuperKExpress_Backlinks_Audit.xlsx")
TPL = os.path.join(REF, "KandK_sample.xlsx")

DIFFS = []
def diff(where, got, exp):
    if got != exp:
        DIFFS.append(f"{where}: got {got!r}, expected {exp!r}")

def font_sig(c):
    f = c.font
    return (f.name, float(f.size) if f.size else None, bool(f.bold),
            (f.color.rgb if f.color and f.color.rgb and isinstance(f.color.rgb, str) else None))

def fill_sig(c):
    pf = c.fill
    fg = pf.fgColor.rgb if pf and pf.fgColor and isinstance(pf.fgColor.rgb, str) else None
    return (pf.patternType, fg)

def border_sig(c):
    b = c.border
    return tuple((s.style if s and s.style else None) for s in (b.left, b.right, b.top, b.bottom))

def align_sig(c):
    a = c.alignment
    return (a.horizontal, a.vertical, bool(a.wrap_text))

DATA_TABS = ["Referring Domain Audit","All Backlinks","Anchor Analysis",
             "Exact Match Anchors","Pending Disavow Lines","Disavow List"]

def main():
    g = openpyxl.load_workbook(GEN)
    t = openpyxl.load_workbook(TPL)
    # tabs + order
    diff("sheet order", g.sheetnames, t.sheetnames)

    for name in DATA_TABS:
        gw, tw = g[name], t[name]
        ncol = tw.max_column
        # headers
        gh = [gw.cell(1,c).value for c in range(1,ncol+1)]
        th = [tw.cell(1,c).value for c in range(1,ncol+1)]
        diff(f"[{name}] headers", gh, th)
        # column widths
        from openpyxl.utils import get_column_letter
        for c in range(1,ncol+1):
            L=get_column_letter(c)
            gwid = round(gw.column_dimensions[L].width or 0,2)
            twid = round(tw.column_dimensions[L].width or 0,2)
            diff(f"[{name}] width col {L}", gwid, twid)
        # header row height
        diff(f"[{name}] header row height", gw.row_dimensions[1].height, tw.row_dimensions[1].height)
        # data row height (row 2)
        diff(f"[{name}] data row height", gw.row_dimensions[2].height, tw.row_dimensions[2].height)
        # freeze
        diff(f"[{name}] freeze", gw.freeze_panes, tw.freeze_panes)
        # autofilter present + anchored at A1 (normalise $ and compare the anchor only;
        # the range row-count legitimately differs by client)
        def filt_anchor(ref):
            r = (ref or "").replace("$", "")
            return r.split(":")[0] if r else ""
        diff(f"[{name}] autofilter anchor", filt_anchor(gw.auto_filter.ref), filt_anchor(tw.auto_filter.ref))
        # gridlines
        diff(f"[{name}] gridlines", gw.sheet_view.showGridLines, tw.sheet_view.showGridLines)
        # header cell styling
        diff(f"[{name}] header font", font_sig(gw.cell(1,1)), font_sig(tw.cell(1,1)))
        diff(f"[{name}] header fill", fill_sig(gw.cell(1,1)), fill_sig(tw.cell(1,1)))
        diff(f"[{name}] header border", border_sig(gw.cell(1,1)), border_sig(tw.cell(1,1)))
        diff(f"[{name}] header align", align_sig(gw.cell(1,1)), align_sig(tw.cell(1,1)))
        # data cell styling (row 2 col 1), if both have data
        if gw.max_row>=2 and tw.max_row>=2:
            diff(f"[{name}] data font", font_sig(gw.cell(2,1)), font_sig(tw.cell(2,1)))
            diff(f"[{name}] data fill", fill_sig(gw.cell(2,1)), fill_sig(tw.cell(2,1)))
            diff(f"[{name}] data border", border_sig(gw.cell(2,1)), border_sig(tw.cell(2,1)))
            diff(f"[{name}] data align", align_sig(gw.cell(2,1)), align_sig(tw.cell(2,1)))

    # number formats (date + percent) on representative cells
    diff("[Referring Domain Audit] First Seen numfmt",
         g["Referring Domain Audit"].cell(2,11).number_format,
         t["Referring Domain Audit"].cell(2,11).number_format)
    diff("[Anchor Analysis] share numfmt",
         g["Anchor Analysis"].cell(2,4).number_format,
         t["Anchor Analysis"].cell(2,4).number_format)

    # Summary
    gs, tsh = g["Summary"], t["Summary"]
    diff("[Summary] share numfmt", gs.cell(4,4).number_format, tsh.cell(4,4).number_format)
    diff("[Summary] gridlines", gs.sheet_view.showGridLines, tsh.sheet_view.showGridLines)
    diff("[Summary] title font", font_sig(gs["B1"]), font_sig(tsh["B1"]))
    from openpyxl.utils import get_column_letter
    for c in range(1,6):
        L=get_column_letter(c)
        diff(f"[Summary] width col {L}", round(gs.column_dimensions[L].width or 0,2),
             round(tsh.column_dimensions[L].width or 0,2))
    diff("[Summary] has logo image", len(getattr(gs,'_images',[]))>0, len(getattr(tsh,'_images',[]))>0)
    def logo_sig(ws):
        imgs = getattr(ws, '_images', [])
        if not imgs: return None
        a = imgs[0].anchor
        try:
            fr = a._from
            return (fr.col, fr.colOff, fr.row, fr.rowOff, a.ext.cx, a.ext.cy)
        except Exception:
            return ("unsupported-anchor",)
    diff("[Summary] logo anchor+size", logo_sig(gs), logo_sig(tsh))

    if DIFFS:
        print(f"STRUCTURAL DIFFERENCES ({len(DIFFS)}):")
        for d in DIFFS:
            print("  DIFF:", d)
        sys.exit(1)
    print("TEMPLATE MATCH: zero structural differences")

if __name__ == "__main__":
    main()
