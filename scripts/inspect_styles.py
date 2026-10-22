import sys
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1])
def fdesc(c):
    f=c.font; fl=c.fill; al=c.alignment; b=c.border
    fill=fl.fgColor.rgb if fl and fl.patternType else None
    bs=''.join(s[0] for s in ['left','right','top','bottom'] if getattr(b,s) and getattr(b,s).style)
    return (f"font={f.name},{f.size},b={f.bold},col={f.color.rgb if f.color else None} | "
            f"fill={fill} | align=h:{al.horizontal},v:{al.vertical},wrap:{al.wrap_text} | border=[{bs}]")

ws=wb['Summary']
print("=== SUMMARY cells (non-empty) ===")
for r in range(1,29):
    for col in range(1,6):
        c=ws.cell(r,col)
        if c.value is not None:
            print(f"  {c.coordinate}: {c.value!r:55.55} | {fdesc(c)}")
# logo
for img in getattr(ws,'_images',[]):
    a=img.anchor
    try:
        print("LOGO anchor _from col,row,colOff,rowOff:", a._from.col, a._from.row, a._from.colOff, a._from.rowOff)
        print("LOGO ext cx,cy:", getattr(a,'ext',None))
    except Exception as e:
        print("logo anchor:", a, e)
    print("LOGO size px:", img.width, img.height)

for tab in ['Referring Domain Audit','Exact Match Anchors']:
    ws=wb[tab]
    print(f"\n=== {tab}: header row1 + data row2 ===")
    print("  H A1:", fdesc(ws['A1']))
    print("  H C1:", fdesc(ws['C1']))
    print("  D A2:", fdesc(ws['A2']))
    print("  D C2:", fdesc(ws['C2']))
    if tab=='Exact Match Anchors':
        print("  F1 (Partner confirmed hdr):", fdesc(ws['F1']))
        print("  F2 (input cell):", fdesc(ws['F2']))
        print("  G2 (final decision):", fdesc(ws['G2']))
        dvs=ws.data_validations.dataValidation
        print("  data validations:", [(dv.sqref, dv.type, dv.formula1) for dv in dvs])
