#!/usr/bin/env python3
"""Build the kbelectric.online referring-domain audit workbook, matching the K&K template."""
import csv, json, sys, re, datetime as dt
from collections import defaultdict, Counter
from urllib.parse import urlparse
sys.path.insert(0, '/home/user/test/scripts')
import classify as C
import template_spec as T
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DATA='/home/user/test/data'
OUT='/home/user/test/output/KBElectric_Backlinks_Audit.xlsx'
CLIENT_TITLE='KB Electric LLC'
TARGET_ROOT='kbelectric.online'
DISAVOW_FILE='disavow_kbelectric.txt'

# ---------- helpers ----------
MULTI_SUFFIX={'co.in','net.in','org.in','co.uk','org.uk','me.uk','com.au','net.au','org.au',
    'co.nz','co.za','com.br','com.mx','co.jp','or.jp','com.sg','com.my','co.id','com.tr',
    'com.ua','co.il','com.ph','com.pk','com.ng','com.co','web.id'}
# multi-tenant platforms: a disavow must target the SUBDOMAIN, never the shared root,
# so these are treated as effective public suffixes (domain:sub.blogspot.com, not domain:blogspot.com)
PLATFORMS={'blogspot.com','wordpress.com','pages.dev','github.io','wixsite.com','weebly.com',
    'tumblr.com','medium.com','substack.com','netlify.app','vercel.app','web.app','000webhostapp.com',
    'myshopify.com','hashnode.dev','notion.site','gitbook.io','strikingly.com','jimdosite.com',
    'glitch.me','repl.co','firebaseapp.com','azurewebsites.net','herokuapp.com','blogspot.in',
    'blogspot.co.uk','blogspot.ca','blogspot.com.au','wordpress.co','square.site','godaddysites.com',
    'webflow.io','sites.google.com','translate.goog','r.appspot.com','onrender.com'}
def reg_root(host):
    host=(host or '').lower().strip().rstrip('.')
    if host.startswith('www.'): host=host[4:]
    for p in PLATFORMS:
        if host==p: return host
        if host.endswith('.'+p):
            label=host[:-(len(p)+1)].split('.')[-1]
            return f"{label}.{p}" if label else host
    parts=host.split('.')
    if len(parts)<=2: return host
    last2='.'.join(parts[-2:]); last3='.'.join(parts[-3:])
    if last2 in MULTI_SUFFIX and len(parts)>=3: return last3
    return last2
def hostof(u):
    try:
        h=urlparse(u if '://' in u else 'http://'+u).netloc.lower()
        return h[4:] if h.startswith('www.') else h
    except: return ''
def norm_url(u):
    u=(u or '').lower().strip()
    u=re.sub(r'^https?://','',u); u=re.sub(r'^www\.','',u)
    return u.rstrip('/')
def to_date(v):
    if v is None or v=='': return None
    s=str(v)
    if s.isdigit(): return dt.date.fromtimestamp(int(s))
    try: return dt.datetime.strptime(s[:10],'%Y-%m-%d').date()
    except: return None

# ---------- load enrichment ----------
enrich=json.load(open(f'{DATA}/batch_enrich.json'))
def enr(key):
    # try exact host, then registrable root
    e=enrich.get(key)
    if e is None: e=enrich.get(reg_root(key))
    return e or {}

# ---------- load Semrush refdomains (IP + Semrush AS) ----------
sr_ref={}
for r in csv.DictReader(open(f'{DATA}/semrush_refdomains.csv',encoding='utf-8'),delimiter=';'):
    d=hostof(r['domain']) or r['domain'].lower()
    sr_ref[d]={'ip':r.get('ip','') or '','as':r.get('domain_score',''),
               'first':to_date(r.get('first_seen')),'last':to_date(r.get('last_seen'))}

# ---------- build unified per-link records ----------
ah_ip={}   # key(root_name_source) -> ip_source  (fallback IP for Ahrefs-only domains)
links=[]  # dict per link
def add_link(key, source_url, target_url, anchor, dofollow, first_seen, title, tool, link_type):
    at=C.anchor_type(anchor, link_type)
    links.append(dict(key=key, root=reg_root(key), source_url=source_url, target_url=target_url,
        anchor=anchor or '', dofollow=dofollow, first=first_seen, title=title or '',
        tool=tool, link_type='Dofollow' if dofollow else 'Nofollow', atype=at,
        nkey=(norm_url(source_url), norm_url(target_url))))

# Semrush links  (key = source host)
for r in csv.DictReader(open(f'{DATA}/semrush_backlinks.csv',encoding='utf-8'),delimiter=';'):
    key=hostof(r['source_url'])
    if not key: continue
    add_link(key, r['source_url'], r['target_url'], r['anchor'],
             r['nofollow']=='false', to_date(r['first_seen']), r.get('source_title',''),
             'Semrush', '')
# Ahrefs links  (key = root_name_source to match enrichment universe)
aj=json.load(open(f'{DATA}/ahrefs_backlinks.json')); ab=aj['backlinks'] if isinstance(aj,dict) else aj
for r in ab:
    key=(r.get('root_name_source') or r.get('name_source') or '').lower()
    if key.startswith('www.'): key=key[4:]
    if not key: continue
    if r.get('ip_source') and key not in ah_ip: ah_ip[key]=r['ip_source']
    add_link(key, r.get('url_from',''), r.get('url_to',''), r.get('anchor',''),
             bool(r.get('is_dofollow')), to_date(r.get('first_seen')), r.get('title',''),
             'Ahrefs', r.get('link_type',''))

# mark Tool=Both for links present in both sources (same normalized source+target)
bykey=defaultdict(set)
for L in links: bykey[L['nkey']].add(L['tool'])
both=set(k for k,s in bykey.items() if {'Semrush','Ahrefs'}<=s)
# dedupe: collapse identical (nkey, anchor) keeping one row, tool reflects union
seen={}
dedup=[]
for L in links:
    sig=(L['nkey'], L['anchor'].strip().lower())
    tool='Both' if L['nkey'] in both else L['tool']
    if sig in seen:
        seen[sig]['tool']='Both' if (seen[sig]['tool']!=tool or tool=='Both') and {seen[sig]['tool'],tool}>={'Semrush','Ahrefs'} else seen[sig]['tool']
        if L['nkey'] in both: seen[sig]['tool']='Both'
        continue
    L=dict(L); L['tool']=tool
    seen[sig]=L; dedup.append(L)
links=dedup
print(f"unified links: {len(links)}")

# ---------- aggregate per audit-key (host) ----------
doms=defaultdict(lambda: dict(anchors=Counter(), atypes=set(), dofollow=False, nofollow=False,
    nlinks=0, first=None, titles=[], srcs=set(), examples=[], targets=Counter(), spamcopy=False))
for L in links:
    d=doms[L['key']]
    d['anchors'][L['anchor']]+=1
    d['atypes'].add(L['atype'])
    d['nlinks']+=1
    if L['dofollow']: d['dofollow']=True
    else: d['nofollow']=True
    if L['atype']==C.T_SPAM: d['spamcopy']=True
    if L['first'] and (d['first'] is None or L['first']<d['first']): d['first']=L['first']
    d['titles'].append(L['title'])
    if L['tool']=='Both': d['srcs'].update(('Semrush','Ahrefs'))
    else: d['srcs'].add(L['tool'])
    d['targets'][L['target_url']]+=1
    d['examples'].append(L)
def source_label(srcs):
    return 'Both' if {'Semrush','Ahrefs'}<=srcs else (next(iter(srcs)) if srcs else '')
def as_int(v):
    try: return int(round(float(v)))
    except: return v if v not in (None,'') else None

# ---------- classify each domain ----------
status={}
for key,d in doms.items():
    e=enr(key)
    tr=e.get('traffic'); drv=e.get('dr')
    status[key]=C.classify_domain(key, d['atypes'], d['dofollow'], d['spamcopy'],
                                  dr=drv, traffic=tr, titles=d['titles'])

# propagate to links
for L in links: L['status']=status[L['key']]

# ---------- In Disavow + disavow line + covered-by-parent ----------
spam_roots=set()
for key,st in status.items():
    if st==C.S_SPAM_DISAVOW: spam_roots.add(reg_root(key))
indis={}; disline={}
for key,st in status.items():
    if st==C.S_SPAM_DISAVOW:
        root=reg_root(key)
        if key!=root and root in status and status.get(root)==C.S_SPAM_DISAVOW:
            indis[key]='Covered by parent'; disline[key]=''
        else:
            indis[key]='Yes'; disline[key]=f'domain:{root}'
    elif st in (C.S_HELD_PARTNER, C.S_HELD_NOFOLLOW):
        indis[key]='Held'; disline[key]=''
    else:
        indis[key]='No'; disline[key]=''

# ---------- site-type heuristic (Exact Match tab) ----------
def site_type(domain, titles, urls):
    d=domain.lower(); blob=(' '.join(titles)+' '+' '.join(urls)).lower()
    if d in C.LEGIT_DOMAINS: return 'Genuine site'
    if any(w in d or w in blob for w in ['bookmark','pinboard','folkd','digg','/story/','mix.com']):
        if 'bookmark' in d or 'bookmark' in blob: return 'Social bookmarking site'
    if any(w in d for w in ['director','listing','list','local','biz','.info','places','find']) \
       or '/detail/' in blob or '/listing' in blob:
        return 'Low quality web directory'
    if any(w in blob for w in ['/article','/story/','/news','press-release','blog','lifestyle','guest']) \
       or any(w in d for w in ['article','news','blog','lifestyle','post','magz','press']):
        return 'Guest article or content farm'
    return 'Low quality web directory'

# ========== BUILD WORKBOOK ==========
wb=openpyxl.Workbook(); wb.remove(wb.active)
HFILL=PatternFill('solid', fgColor=T.HEADER_FILL)
WFILL=PatternFill('solid', fgColor=T.WHITE_FILL)
def hdr_font(sz=9): return Font(name=T.FONT_NAME, bold=True, color=T.HEADER_FONT, size=sz)
def body_font(sz=9, bold=False): return Font(name=T.FONT_NAME, bold=bold, color=T.BODY_COLOR, size=sz)
CEN=Alignment(horizontal='center', wrap_text=True)
LEFTW=Alignment(wrap_text=True, vertical='top')

def set_widths(ws, name):
    for col,w in T.WIDTHS.get(name,{}).items():
        ws.column_dimensions[col].width=w
def write_header(ws, headers, row=1):
    for ci,h in enumerate(headers,1):
        c=ws.cell(row,ci,h); c.font=hdr_font(); c.fill=HFILL; c.alignment=CEN

# ----- helper to pick example page/title/top anchor -----
def rep(d):
    top_anchor=d['anchors'].most_common(1)[0][0] if d['anchors'] else ''
    ex=None
    for L in d['examples']:
        if L['anchor']==top_anchor: ex=L; break
    ex=ex or (d['examples'][0] if d['examples'] else None)
    return top_anchor, (ex['source_url'] if ex else ''), (ex['title'] if ex else '')

# ===== Tab 2: Referring Domain Audit =====
ws=wb.create_sheet('Referring Domain Audit'); set_widths(ws,'Referring Domain Audit')
write_header(ws, T.TABS['Referring Domain Audit']); ws.freeze_panes='B2'
audit_rows=sorted(doms.items(), key=lambda kv:(-kv[1]['nlinks'], kv[0]))
r=2
for key,d in audit_rows:
    e=enr(key); sr=sr_ref.get(key,{})
    top_anchor, ex_page, ex_title=rep(d)
    as_val=as_int(sr.get('as') or e.get('dr'))
    ip_val=sr.get('ip') or ah_ip.get(key) or ''
    dof=sum(1 for L in d['examples'] if L['dofollow']); nof=d['nlinks']-dof
    tgt_root=reg_root(hostof(d['targets'].most_common(1)[0][0])) if d['targets'] else TARGET_ROOT
    vals=[key, status[key], as_val, d['nlinks'], e.get('refdomains'), e.get('traffic'),
          indis[key], tgt_root or TARGET_ROOT, dof, nof, d['first'], top_anchor, ex_page,
          ex_title, source_label(d['srcs']), ip_val, disline[key]]
    for ci,v in enumerate(vals,1):
        c=ws.cell(r,ci,v); c.font=body_font(); c.fill=WFILL
        if ci in (4,5,6,9,10): c.number_format=T.FMT_INT; c.alignment=Alignment(horizontal='center')
        if ci==11: c.number_format=T.FMT_DATE
        if ci in (13,14): c.alignment=LEFTW
    r+=1
AUDIT_N=r-2

# ===== Tab 3: All Backlinks =====
ws=wb.create_sheet('All Backlinks'); set_widths(ws,'All Backlinks')
write_header(ws, T.TABS['All Backlinks']); ws.freeze_panes='B2'
r=2
for L in sorted(links, key=lambda x:(x['key'], x['anchor'])):
    vals=[L['key'], L['status'], L['tool'], L['anchor'], L['atype'], L['link_type'],
          L['source_url'], L['title'], L['target_url']]
    for ci,v in enumerate(vals,1):
        c=ws.cell(r,ci,v); c.font=body_font(); c.fill=WFILL
        if ci in (4,7,8,9): c.alignment=LEFTW
    r+=1
ALLBL_N=r-2

# ===== Tab 4: Anchor Analysis =====
ws=wb.create_sheet('Anchor Analysis'); set_widths(ws,'Anchor Analysis')
write_header(ws, T.TABS['Anchor Analysis']); ws.freeze_panes='B2'
anchor_agg=defaultdict(lambda: dict(bl=0, doms=set(), spamdoms=set(), atype=None))
for L in links:
    a=anchor_agg[L['anchor']]; a['bl']+=1; a['doms'].add(L['key']); a['atype']=L['atype']
    if L['status'] in (C.S_SPAM_DISAVOW, C.S_SPAM_NOFOLLOW): a['spamdoms'].add(L['key'])
total_bl=len(links)
r=2
for anchor,a in sorted(anchor_agg.items(), key=lambda kv:-kv[1]['bl']):
    vals=[anchor, a['atype'], a['bl'], a['bl']/total_bl, len(a['doms']), len(a['spamdoms'])]
    for ci,v in enumerate(vals,1):
        c=ws.cell(r,ci,v); c.font=body_font(); c.fill=WFILL
        if ci==1: c.alignment=LEFTW
        if ci in (3,5,6): c.number_format=T.FMT_INT; c.alignment=Alignment(horizontal='center')
        if ci==4: c.number_format=T.FMT_PCT; c.alignment=Alignment(horizontal='center')
    r+=1
ANCHOR_N=r-2

# ===== Tab 5: Exact Match Anchors =====
ws=wb.create_sheet('Exact Match Anchors'); set_widths(ws,'Exact Match Anchors')
write_header(ws, T.TABS['Exact Match Anchors']); ws.freeze_panes='B2'
em_seen=set(); em_rows=[]
for L in links:
    if L['atype'] in (C.T_EXACT, C.T_EXACT_GEO):
        sig=(L['key'], L['anchor'].strip().lower())
        if sig in em_seen: continue
        em_seen.add(sig)
        market='Oklahoma' if L['atype']==C.T_EXACT_GEO else 'Non-geo'
        d=doms[L['key']]
        stype=site_type(L['key'], d['titles'], [x['source_url'] for x in d['examples']])
        em_rows.append([L['key'], L['source_url'], L['anchor'], market, stype, None, None])
em_rows.sort(key=lambda x:x[0])
r=2
for vals in em_rows:
    for ci,v in enumerate(vals,1):
        c=ws.cell(r,ci,v); c.font=body_font(); c.fill=WFILL
        if ci in (2,3): c.alignment=LEFTW
    r+=1
EM_N=r-2

# ===== Tab 6: Pending Disavow Lines =====
ws=wb.create_sheet('Pending Disavow Lines'); set_widths(ws,'Pending Disavow Lines')
write_header(ws, T.TABS['Pending Disavow Lines']); ws.freeze_panes='B2'
rec=f'Append only if the partner does not claim it, to {DISAVOW_FILE}'
pending=sorted(reg_root(k) for k,st in status.items() if st==C.S_HELD_PARTNER)
pending=sorted(set(pending))
r=2
for root in pending:
    vals=[root, f'domain:{root}', rec]
    for ci,v in enumerate(vals,1):
        c=ws.cell(r,ci,v); c.font=body_font(); c.fill=WFILL
        if ci==3: c.alignment=LEFTW
    r+=1
PEND_N=r-2

# ===== Tab 7: Disavow List =====
ws=wb.create_sheet('Disavow List'); set_widths(ws,'Disavow List')
write_header(ws, T.TABS['Disavow List']); ws.freeze_panes='B2'
# aggregate spam-disavow by registrable root
dis_agg=defaultdict(lambda: dict(bl=0, anchors=Counter(), as_v=None, rd=None, tr=None, kw=None))
for key,st in status.items():
    if indis.get(key)=='Yes':
        root=reg_root(key); e=enr(key); sr=sr_ref.get(key,{}); d=doms[key]
        a=dis_agg[root]; a['bl']+=d['nlinks']
        for anc,n in d['anchors'].items(): a['anchors'][anc]+=n
        asv=sr.get('as') or e.get('dr')
        if asv not in (None,''): a['as_v']=max(a['as_v'] or 0, float(asv))
        if e.get('refdomains') is not None: a['rd']=max(a['rd'] or 0, e['refdomains'])
        if e.get('traffic') is not None: a['tr']=max(a['tr'] or 0, e['traffic'])
        if e.get('org_keywords') is not None: a['kw']=max(a['kw'] or 0, e['org_keywords'])
r=2
for root in sorted(dis_agg):
    a=dis_agg[root]; top=a['anchors'].most_common(1)[0][0] if a['anchors'] else ''
    vals=[f'domain:{root}', as_int(a['as_v']), a['bl'], a['rd'], a['tr'], a['kw'], top]
    for ci,v in enumerate(vals,1):
        c=ws.cell(r,ci,v); c.font=body_font(); c.fill=WFILL
        if ci in (2,3,4,5,6): c.number_format=T.FMT_INT; c.alignment=Alignment(horizontal='center')
        if ci==7: c.alignment=LEFTW
    r+=1
DIS_N=r-2

# ===== Tab 1: Summary =====
ws=wb.create_sheet('Summary', 0); set_widths(ws,'Summary')
ws.merge_cells('B1:E2')
ws['B1']=CLIENT_TITLE; ws['B1'].font=Font(name=T.FONT_NAME,bold=True,size=21,color=T.TITLE_COLOR)
ws['B1'].alignment=Alignment(horizontal='center', vertical='center')
# status table
status_dom=Counter(status.values())
status_bl=Counter()
for L in links: status_bl[L['status']]+=1
total_dom=sum(status_dom.values())
write_header(ws, ['Status','Domains','Backlinks','Share of Domains','What it means'], row=3)
rr=4
for st in T.STATUS_ORDER:
    dn=status_dom.get(st,0); bn=status_bl.get(st,0)
    ws.cell(rr,1,st).font=body_font()
    ws.cell(rr,2,dn).font=body_font(); ws.cell(rr,2).number_format=T.FMT_INT; ws.cell(rr,2).alignment=Alignment(horizontal='center')
    ws.cell(rr,3,bn).font=body_font(); ws.cell(rr,3).number_format=T.FMT_INT; ws.cell(rr,3).alignment=Alignment(horizontal='center')
    ws.cell(rr,4,(dn/total_dom if total_dom else 0)).font=body_font(); ws.cell(rr,4).number_format=T.FMT_PCT; ws.cell(rr,4).alignment=Alignment(horizontal='center')
    ws.cell(rr,5,T.STATUS_MEANING[st]).font=body_font(); ws.cell(rr,5).alignment=LEFTW
    for ci in range(1,6): ws.cell(rr,ci).fill=WFILL
    rr+=1
# total row
ws.cell(rr,1,'Total').font=body_font(bold=True)
ws.cell(rr,2,total_dom).font=body_font(bold=True); ws.cell(rr,2).number_format=T.FMT_INT; ws.cell(rr,2).alignment=Alignment(horizontal='center')
ws.cell(rr,3,len(links)).font=body_font(bold=True); ws.cell(rr,3).number_format=T.FMT_INT; ws.cell(rr,3).alignment=Alignment(horizontal='center')
ws.cell(rr,4,1.0).font=body_font(bold=True); ws.cell(rr,4).number_format=T.FMT_PCT; ws.cell(rr,4).alignment=Alignment(horizontal='center')
total_row=rr; rr+=1   # Key findings starts immediately after Total (row 11), matching template
# key findings
spam_dom=status_dom.get(C.S_SPAM_DISAVOW,0)+status_dom.get(C.S_SPAM_NOFOLLOW,0)
held_dom=status_dom.get(C.S_HELD_PARTNER,0)+status_dom.get(C.S_HELD_NOFOLLOW,0)
disavow_n=DIS_N
# anchor profile
at_bl=Counter(L['atype'] for L in links)
at_dom=defaultdict(set)
for L in links: at_dom[L['atype']].add(L['key'])
top_spam_anchor=max((a for a in anchor_agg.items() if a[1]['atype']==C.T_SPAM), key=lambda kv:len(kv[1]['doms']), default=(None,{'doms':set()}))
# home vs service page split by anchor group
def is_home(u):
    p=urlparse(u).path.rstrip('/'); return p=='' or p=='/'
spam_home=sum(1 for L in links if L['atype']==C.T_SPAM and is_home(L['target_url']))
spam_tot=at_bl.get(C.T_SPAM,0)
comm_links=[L for L in links if L['atype'] in (C.T_EXACT,C.T_EXACT_GEO)]
comm_service=sum(1 for L in comm_links if not is_home(L['target_url']))
# timing bursts
def ym(d): return (d.year,d.month) if d else None
allm=Counter(ym(L['first']) for L in links if L['first'])
peak=max(allm.items(), key=lambda kv:kv[1]) if allm else (None,0)
heldm=Counter(ym(L['first']) for L in links if L['status'] in (C.S_HELD_PARTNER,C.S_HELD_NOFOLLOW) and L['first'])
hpeak=max(heldm.items(), key=lambda kv:kv[1]) if heldm else (None,0)
MON=['','Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
def mlabel(t): return f"{MON[t[1]]} {t[0]}" if t and t[0] else "a single month"
def mlab(item): return mlabel(item[0]) if item and item[0] else "a single month"
held_notraffic=sum(1 for k,st in status.items() if st in (C.S_HELD_PARTNER,C.S_HELD_NOFOLLOW)
                   and not (enr(k).get('traffic') or 0))
spam_anchor_txt=(top_spam_anchor[0] or '')[:60]
findings=[
 f"{spam_dom} of the {total_dom} referring domains are spam. {status_dom.get(C.S_SPAM_DISAVOW,0)} carry at least one dofollow link and go in the disavow file ({disavow_n} unique domain lines); the other {status_dom.get(C.S_SPAM_NOFOLLOW,0)} are nofollow and pass no value, so they need no action.",
 f"Two patterns, easy to tell apart. The link scheme is a keyword-stuffed SEO sales-copy flood: {'all ' if spam_home==spam_tot else str(spam_home)+' of the '}{spam_tot} sales-copy links point at the home page. The {len(comm_links)} exact-match commercial links behave differently — {comm_service} of them hit KB's service pages (sports / stadium / court lighting, commercial electrical).",
 f"The single biggest spam anchor, on {len(top_spam_anchor[1]['doms'])} domains, is backlink sales copy that names the client's own domain — e.g. “{spam_anchor_txt}…”. An anchor advertising a link-selling service, not the business, is the clearest sign of a bought/PBN link.",
 f"The attack is recent and concentrated: the biggest single month is {mlab(peak)} with {peak[1]} links first seen — many domains stood up in a short window, which is the signature of an automated network rather than organic growth.",
 f"{held_dom} domains are held for the previous partner to confirm, in case any were commissioned. They use real commercial anchors (not sales copy) and cluster around {mlab(hpeak)} ({hpeak[1]} links), consistent with a deliberate build. They are NOT in the disavow file until the partner confirms, so nothing paid-for is discarded by accident.",
 f"Held does not mean good: {held_notraffic} of the {held_dom} held domains have no organic traffic at all. They are kept out of the disavow file only because the anchor could indicate a paid placement — confirm, then disavow the unclaimed ones using Tab 'Pending Disavow Lines'.",
 f"Unlike some networks, these domains are spread across many IP subnets (no single hosting cluster), so IP is not a reliable filter here — the anchor text and the sales-copy pattern are what separate the attack from the commissioned links.",
]
ws.cell(rr,1,'Key findings').font=Font(name=T.FONT_NAME,bold=True,size=10,color=T.BODY_COLOR); rr+=1
for f in findings:
    c=ws.cell(rr,1,f); c.font=body_font(); c.alignment=LEFTW
    ws.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=5)  # bullet spans A:E like template
    rr+=1
ws.cell(rr,1,'Anchor text profile (all backlinks)').font=Font(name=T.FONT_NAME,bold=True,size=11,color=T.TITLE_COLOR); rr+=1
write_header(ws, ['Anchor type','Backlinks','Share of Backlinks','Referring Domains'], row=rr); rr+=1
for atype in T.ANCHOR_TYPES:
    bn=at_bl.get(atype,0)
    ws.cell(rr,1,atype).font=body_font()
    ws.cell(rr,2,bn).font=body_font(); ws.cell(rr,2).number_format=T.FMT_INT; ws.cell(rr,2).alignment=Alignment(horizontal='center')
    ws.cell(rr,3,(bn/total_bl if total_bl else 0)).font=body_font(); ws.cell(rr,3).number_format=T.FMT_PCT; ws.cell(rr,3).alignment=Alignment(horizontal='center')
    ws.cell(rr,4,len(at_dom.get(atype,set()))).font=body_font(); ws.cell(rr,4).number_format=T.FMT_INT; ws.cell(rr,4).alignment=Alignment(horizontal='center')
    for ci in range(1,5): ws.cell(rr,ci).fill=WFILL
    rr+=1

wb.save(OUT)
print("SAVED", OUT)
print(f"Audit domains={AUDIT_N} AllBacklinks={ALLBL_N} Anchors={ANCHOR_N} ExactMatch={EM_N} Pending={PEND_N} Disavow={DIS_N}")
print("Status(domains):", dict(status_dom))
print("Status(backlinks):", dict(status_bl))
print("AnchorType(backlinks):", dict(at_bl))
