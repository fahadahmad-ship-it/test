#!/usr/bin/env python3
# Build Harbinger-format deliverables for scapersllc.com.
# python3 -I scapers/scripts/build.py <classified.csv> <anchors.csv> <ahrefs_backlinks.csv> <ahrefs_refdomains.csv> <history.csv> <outdir>
import sys, csv, json, datetime, collections, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
sys.path.insert(0, os.path.join(os.path.dirname(__file__)))
import classify as C

CLIENT_NAME = "Scapers, LLC"
CDOM = "scapersllc.com"
ACCENT = "FF0C343D"; WHITE = "FFFFFFFF"
LOGO = os.path.join(os.path.dirname(__file__), "..", "data", "harbinger_logo.png")
csv.field_size_limit(10**7)

def F(sz, b=False, col="FF000000"): return Font(name="Geist", size=sz, bold=b, color=col)
accent_fill=PatternFill("solid",fgColor=ACCENT); white_fill=PatternFill("solid",fgColor=WHITE)
thin=Side(style="thin",color="FF000000")
hdr_border=Border(left=thin,right=thin,top=thin); full_border=Border(left=thin,right=thin,top=thin,bottom=thin); no_border=Border()
c_cen=Alignment(horizontal="center",vertical="center",wrap_text=True)
c_cennw=Alignment(horizontal="center",vertical="center")
c_left=Alignment(horizontal="left",vertical="center"); c_vc=Alignment(vertical="center"); c_vcw=Alignment(vertical="center",wrap_text=True)
FNT_TITLE=F(21,True,ACCENT); FNT_HDR=F(9,True,WHITE); FNT_HDR10=F(10,True,WHITE)
FNT_DATA=F(9,False,"FF000000"); FNT_SUMB=F(9,True,"FF000000"); FNT_SECT=F(11,True,ACCENT); FNT_KEY=F(10,True,"FF000000"); FNT_INPUT=F(9,False,WHITE)

def ts(x):
    try: return datetime.datetime.utcfromtimestamp(int(x))
    except: return None

def hdr(ws,r,c,v):
    cell=ws.cell(r,c,v); cell.font=FNT_HDR; cell.fill=accent_fill; cell.border=hdr_border; cell.alignment=c_cen; return cell
def dcell(ws,r,c,v,num=None,al=c_vc):
    cell=ws.cell(r,c,v); cell.font=FNT_DATA; cell.fill=white_fill; cell.border=no_border; cell.alignment=al
    if num: cell.number_format=num
    return cell
def style_sheet(ws,ncols,nrows,rowht,widths):
    ws.freeze_panes="B2"; ws.sheet_view.showGridLines=True
    ws.auto_filter.ref=f"A1:{get_column_letter(ncols)}{max(nrows,1)}"
    ws.row_dimensions[1].height=31.5
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for rr in range(2,nrows+1): ws.row_dimensions[rr].height=rowht

GAMBLING_EV=("gambling","off-topic","betsat","maxwin","sahabet","betwoon","mobilbahis","grandpashabet","betpark","casino","flood host")
def reason_group(ev):
    e=ev.lower()
    if "hacked" in e or "compromised" in e: return "Hacked third-party site (gambling injection)"
    if any(k in e for k in GAMBLING_EV): return "Off-topic gambling anchors (July 2026 flood)"
    if "link-selling" in e or "pbn sales-copy" in e: return "Link-selling / PBN sales-copy network"
    if "host-name signature" in e: return "Spam host-name signature (SEO-tool / directory / PBN host)"
    if "link-scheme injection" in e: return "Naked-URL link-scheme injection"
    return "Other spam (evidence recorded per domain)"

def main():
    cls,anch_p,ah_p,ahrd_p,hist_p,outdir=sys.argv[1:7]
    os.makedirs(outdir,exist_ok=True)
    rows=[]
    for r in csv.DictReader(open(cls,encoding="utf-8")):
        r["AS"]=int(r["AS"] or 0); r["backlinks"]=int(r["backlinks"] or 0); r["first_seen"]=int(r["first_seen"] or 0)
        rows.append(r)
    by_dom={r["domain"]:r for r in rows}
    # Metrics are Semrush-only. Semrush holds no organic traffic/keyword data for these spam domains
    # (domain_rank returns NOTHING FOUND = a real zero), so Traffic and Organic Keywords are recorded as 0.
    # Ahrefs is used ONLY for link detail (anchors, targets, dofollow, example pages) — never a metric.
    for r in rows: r["traffic"]=0

    # ---- tallies ----
    stat_order=["Spam (Disavow)","Held (Pending Partner)","Held (Nofollow, No Action)","Spam (Nofollow, No Action)","Keep","Own Site"]
    dc=collections.Counter(r["status"] for r in rows); bc=collections.Counter()
    for r in rows: bc[r["status"]]+=r["backlinks"]
    total_d=len(rows); total_b=sum(r["backlinks"] for r in rows)
    spam_file=dc.get("Spam (Disavow)",0); spam_nf=dc.get("Spam (Nofollow, No Action)",0)
    held=dc.get("Held (Pending Partner)",0)+dc.get("Held (Nofollow, No Action)",0)
    held_low_as=sum(1 for r in rows if r["status"].startswith("Held") and r["AS"]<=5)

    # ---- anchors (Semrush) ----
    anchors=[]
    if os.path.exists(anch_p):
        rr=csv.reader(open(anch_p,encoding="utf-8"),delimiter=";"); next(rr,None)
        for x in rr:
            if len(x)>=3:
                try: anchors.append((x[0],int(x[1] or 0),int(x[2] or 0)))
                except: pass
    tot_anchor_bl=sum(b for _,_,b in anchors) or 1
    at_b=collections.Counter()
    for a,dn,bn in anchors: at_b[C.anchor_type(a)]+=bn
    # distinct referring domains per anchor-type (from per-domain anchor sets; domains with no anchor data are skipped)
    perdom_path=os.path.join(os.path.dirname(cls),"perdomain.json")
    perdom=json.load(open(perdom_path)) if os.path.exists(perdom_path) else {}
    at_dom=collections.Counter()
    for dom,pd in perdom.items():
        types={C.anchor_type(a) for a,_ in pd.get("anchors",[]) if a}
        for ty in types: at_dom[ty]+=1

    # ================= WORKBOOK =================
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Summary"; ws.sheet_view.showGridLines=False
    for i,w in enumerate([38.86,16.29,17.57,18.71,120.0],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.merge_cells("B1:E2"); t=ws.cell(1,2,CLIENT_NAME); t.font=FNT_TITLE; t.alignment=c_cennw
    try:
        im=XLImage(LOGO); ws.add_image(im,"A1")
    except Exception: pass
    means={"Spam (Disavow)":"Confirmed spam with at least one dofollow link. Included in the disavow file.",
           "Held (Pending Partner)":"Exact-match commercial/geo anchor. May have been commissioned. Not in the file until the supplier (previous SEO partner) confirms.",
           "Held (Nofollow, No Action)":"Same, but every link is nofollow so it passes no value. No disavow needed.",
           "Spam (Nofollow, No Action)":"Spam, but every link is nofollow and passes no value, so no disavow is needed.",
           "Keep":"Genuine or harmless link. Keep.","Own Site":"Client-owned domain / redirect."}
    r0=3
    for c,h in enumerate(["Status","Domains","Backlinks","Share of Domains","What it means"],1):
        cell=ws.cell(r0,c,h); cell.font=FNT_HDR10; cell.fill=accent_fill; cell.border=full_border; cell.alignment=c_cen
    rr=r0+1
    present=[s for s in stat_order if dc.get(s,0)>0]
    for s in present:
        vals=[s,dc[s],bc[s],dc[s]/total_d,means[s]]
        for c,v in enumerate(vals,1):
            cell=ws.cell(rr,c,v); cell.font=FNT_SUMB; cell.border=full_border
            cell.alignment=c_left if c in (1,5) else c_cennw
            if c==3: cell.number_format='#,##0'
            if c==4: cell.number_format='0.0%'
        rr+=1
    for c,v in enumerate(["Total",total_d,total_b,1.0,""],1):
        cell=ws.cell(rr,c,v); cell.font=FNT_SUMB; cell.border=full_border; cell.alignment=c_left if c in (1,5) else c_cennw
        if c==3: cell.number_format='#,##0'
        if c==4: cell.number_format='0.0%'
    # key findings
    kr=rr+2; ws.cell(kr,1,"Key findings").font=FNT_KEY
    findings=[
      f"{spam_file+spam_nf:,} of the {total_d:,} referring domains are spam. {spam_file:,} have a dofollow link and go in the disavow file; {spam_nf:,} are nofollow-only and need no action. Semrush shows 26,912 dofollow vs 1,158 nofollow links overall.",
      f"The attack has exploded since the brief. The brief's snapshot was 17,457 links / 1,119 referring domains (~July 2026). The live profile is {total_b:,} backlinks / {total_d:,} referring domains, and Semrush records 4,961 domains as of 1 Oct 2026 (up from 1,365 on 1 Sep).",
      f"Two separate floods. July 2026: ~13,231 off-topic Turkish-gambling links (betsat, maxwin, sahabet, betwoon) concentrated on a few high-volume hosts (piggington.com alone carries 8,636). Sep-Oct 2026: ~3,600 fresh PBN domains carrying link-selling sales copy.",
      f"One file, for {CDOM}. All spam source domains are disavowed on this one property. Some flood links show a redirect_url to sister dumpster brands (stldumpsters.com, stldumpster.com, renthisdumpster.com, stlrolloffs.com, farmingtondumpsters.com); if the client controls those Search Console properties they each need their own file.",
      "Two campaigns, told apart by target. The gambling and PBN sales-copy links point at the home page; the exact-match 'dumpster rental <MO city>' anchors point at deep service/product pages. The home-page flood is the attack.",
      "The dominant anchor is link-sales copy that names the client and its sister brands, e.g. 'farmingtondumpsters.com premium authority backlinks for better search visibility' (1,474 domains) and 'High Quality Dofollow Backlinks DA 50 PA 40 Premium PBN Network Service ... buy backlinks online cheap'.",
      f"{held:,} domains are HELD for the supplier (previous SEO partner) to confirm, in case any exact-match 'dumpster rental <city> MO' placements were commissioned. Held is not the same as valuable: {held_low_as:,} of the {held:,} carry a Semrush Authority Score of 5 or below and no organic data. They are excluded from the file pending that conversation.",
    ]
    fr=kr+1
    for fnd in findings:
        ws.merge_cells(start_row=fr,start_column=1,end_row=fr,end_column=5)
        cell=ws.cell(fr,1,fnd); cell.font=F(9); cell.alignment=c_vcw; ws.row_dimensions[fr].height=22.5; fr+=1
    ar=fr+1; ws.cell(ar,1,"Anchor text profile (all backlinks)").font=FNT_SECT; ar+=1
    for c,h in enumerate(["Anchor type","Backlinks","Share of Backlinks","Referring Domains"],1):
        cell=ws.cell(ar,c,h); cell.font=FNT_HDR10; cell.fill=accent_fill; cell.border=full_border; cell.alignment=c_cen
    ar+=1
    for atype,bl in sorted(at_b.items(),key=lambda x:-x[1]):
        if bl==0: continue
        for c,v in enumerate([atype,bl,bl/tot_anchor_bl,at_dom.get(atype,0)],1):
            cell=ws.cell(ar,c,v); cell.font=FNT_SUMB; cell.border=full_border; cell.alignment=c_left if c==1 else c_cennw
            if c==2: cell.number_format='#,##0'
            if c==3: cell.number_format='0.0%'
        ar+=1
    for rrr in range(1,ar+1):
        if ws.row_dimensions[rrr].height is None: ws.row_dimensions[rrr].height=22.5

    # --- Tab 2 Referring Domain Audit ---
    ws=wb.create_sheet("Referring Domain Audit")
    H=["Domain","Status","AS","Backlinks","RD's","Traffic","In Disavow","Links Point To","Dofollow Links","Nofollow Links","First Seen","Top Anchor Text","Example Linking Page","Linking Page Title","Source","IP","Disavow Line"]
    for c,h in enumerate(H,1): hdr(ws,1,c,h)
    r=2
    for row in sorted(rows,key=lambda x:(stat_order.index(x["status"]) if x["status"] in stat_order else 9,-x["backlinks"])):
        indis="Yes" if row["status"]=="Spam (Disavow)" else ("Held" if row["status"].startswith("Held") else "No")
        dl=f"domain:{row['host']}" if row["status"]=="Spam (Disavow)" else ""
        fsd=ts(row["first_seen"]); src="Semrush+Ahrefs" if row.get("top_anchor") else "Semrush"
        vals=[row["domain"],row["status"],row["AS"],row["backlinks"],"",row["traffic"],indis,row["target"],
              row["dofollow"],row["nofollow"],fsd,row["top_anchor"],row["example_url"],row["example_title"],src,row["ip"],dl]
        for c,v in enumerate(vals,1):
            cell=dcell(ws,r,c,v)
            if c in (3,4,6,9,10): cell.number_format='#,##0'; cell.alignment=c_cennw
            if c==11 and fsd: cell.number_format='d mmm yyyy'; cell.alignment=c_cennw
        r+=1
    style_sheet(ws,17,r-1,15.75,[36,24.71,9.57,13,9.43,9.71,12.43,21.29,14,8.71,12,45,8.71,40,14,15,14])

    # --- Tab 3 All Backlinks (Ahrefs link-level) ---
    ws=wb.create_sheet("All Backlinks")
    for c,h in enumerate(["Referring Domain","Status","Tool","Anchor Text","Anchor Type","Link Type","Linking Page URL","Linking Page Title","Target URL"],1): hdr(ws,1,c,h)
    r=2
    if os.path.exists(ah_p):
        f=open(ah_p,encoding="utf-8",errors="replace",newline=""); delim=';' if ';' in f.readline() else ','; f.seek(0)
        rr=csv.DictReader(f,delimiter=delim); cm={c.lower():c for c in rr.fieldnames}
        def g(row,*n):
            for x in n:
                if x in cm: return row.get(cm[x],"")
            return ""
        for row in rr:
            root=(g(row,"root_name_source") or g(row,"name_source")).strip().lower()
            st=by_dom.get(root,{}).get("status","(Ahrefs-only, not in Semrush set)")
            anc=g(row,"anchor"); lt="Dofollow" if str(g(row,"is_dofollow")).lower() in ("true","1") else "Nofollow"
            for c,v in enumerate([root,st,"Ahrefs",anc,C.anchor_type(anc),lt,g(row,"url_from"),g(row,"title"),g(row,"url_to")],1): dcell(ws,r,c,v)
            r+=1
    style_sheet(ws,9,r-1,22.5,[36,24.71,9.57,55,32,10,55,53.14,30.57])

    # --- Tab 4 Anchor Analysis ---
    ws=wb.create_sheet("Anchor Analysis")
    for c,h in enumerate(["Anchor Text","Anchor Type","Backlinks","Share of Backlinks","Referring Domains","Spam Domains Using It"],1): hdr(ws,1,c,h)
    r=2
    for a,dn,bn in sorted(anchors,key=lambda x:-x[2]):
        at=C.anchor_type(a); sd=dn if at.startswith("Spam") else 0
        for c,v in enumerate([a,at,bn,bn/tot_anchor_bl,dn,sd],1):
            cell=dcell(ws,r,c,v)
            if c in (3,5,6): cell.number_format='#,##0'; cell.alignment=c_cennw
            if c==4: cell.number_format='0.0%'; cell.alignment=c_cennw
        r+=1
    style_sheet(ws,6,r-1,22.5,[115.86,37.71,14,8.71,12,20.29])

    # --- Tab 5 Exact Match Anchors (Held) ---
    ws=wb.create_sheet("Exact Match Anchors")
    for c,h in enumerate(["Referring domain","Example backlink (linking page)","Exact-match commercial anchor","Market","Site type (indicative)","Partner confirmed? (Y/N)","Final decision"],1): hdr(ws,1,c,h)
    r=2
    held_rows=[x for x in rows if x["status"].startswith("Held")]
    for row in sorted(held_rows,key=lambda x:x["domain"]):
        anc=row.get("top_anchor",""); at=C.anchor_type(anc)
        market="Geo (Missouri)" if at=="Exact match commercial (geo)" else ("Non-geo" if at=="Exact match commercial" else "Brand / naked URL")
        stype="Sister dumpster brand" if row["domain"] in C.SISTER_DUMPSTER else "Low-quality web property / directory"
        for c,v in enumerate([row["domain"],row.get("example_url",""),anc,market,stype,"",""],1):
            cell=dcell(ws,r,c,v)
            if c in (6,7): cell.fill=accent_fill; cell.font=FNT_INPUT
        r+=1
    style_sheet(ws,7,r-1,22.5,[25.43,59,53.14,16,24.57,15.43,15])

    # --- Tab 6 Pending Disavow Lines ---
    ws=wb.create_sheet("Pending Disavow Lines")
    for c,h in enumerate(["Referring domain","Disavow line (append to disavow.txt once confirmed)","Recommendation"],1): hdr(ws,1,c,h)
    r=2
    for row in sorted(held_rows,key=lambda x:x["domain"]):
        for c,v in enumerate([row["domain"],f"domain:{row['host']}",f"Append only if the supplier does not claim it, to disavow_{CDOM}.txt"],1): dcell(ws,r,c,v)
        r+=1
    style_sheet(ws,3,r-1,22.5,[32,44,48])

    # --- Tab 7 Disavow List ---
    ws=wb.create_sheet("Disavow List")
    for c,h in enumerate(["Disavow Line","AS","Backlinks","RD's","Traffic","Organic Keywords","Top Anchor Text"],1): hdr(ws,1,c,h)
    disavow=dedupe_disavow([x for x in rows if x["status"]=="Spam (Disavow)"])
    r=2
    for row in sorted(disavow,key=lambda x:x["host"]):
        for c,v in enumerate([f"domain:{row['host']}",row["AS"],row["backlinks"],"",row["traffic"],0,row.get("top_anchor","")],1):
            cell=dcell(ws,r,c,v)
            if c in (2,3,5,6): cell.number_format='#,##0'; cell.alignment=c_cennw
        r+=1
    style_sheet(ws,7,r-1,22.5,[44,14,15.57,14,11.14,14.86,98.86])

    xlsx=os.path.join(outdir,"Scapers_LLC_Backlinks_Audit.xlsx"); wb.save(xlsx)

    # ================= DISAVOW FILE =================
    groups=collections.defaultdict(list)
    for row in disavow: groups[reason_group(row["evidence"])].append(row)
    order=["Off-topic gambling anchors (July 2026 flood)","Hacked third-party site (gambling injection)",
           "Link-selling / PBN sales-copy network","Spam host-name signature (SEO-tool / directory / PBN host)",
           "Naked-URL link-scheme injection","Other spam (evidence recorded per domain)"]
    lines=[]
    lines.append(f"# Google disavow file for {CDOM}")
    lines.append(f"# Generated {datetime.date.today().isoformat()} by Harbinger Marketing from a full live-profile backlink audit.")
    lines.append(f"# Search Console property: {CDOM}  (one-time disavow)")
    lines.append(f"# Total disavowed referring domains: {len(disavow)}")
    lines.append("#")
    lines.append("# IMPORTANT — do NOT upload before the supplier (previous SEO partner) conversation.")
    lines.append(f"#   {held} domains carrying exact-match 'dumpster rental <MO city>' anchors are HELD OUT of this file,")
    lines.append("#   pending confirmation of who commissioned them. 'Supplier' = the role K&K's brief calls 'supplier';")
    lines.append("#   in this audit it is Scapers' previous SEO partner. Those held lines are in the workbook (Pending Disavow Lines tab).")
    lines.append("# If a disavow file already exists on this property, uploading REPLACES it — merge, do not overwrite blindly.")
    lines.append("#   If the client has checked Search Console and there is none, this is the first and only file.")
    lines.append("# Sister dumpster brands (stldumpsters.com, stldumpster.com, stldumpsterental.com, renthisdumpster.com,")
    lines.append("#   stlrolloffs.com, farmingtondumpsters.com) that redirect in need their OWN file per property if the client controls them.")
    lines.append("#")
    for grp in order:
        if grp not in groups: continue
        gl=sorted(groups[grp],key=lambda x:x["host"])
        lines.append("")
        lines.append(f"# === {grp} — {len(gl)} domains ===")
        for row in gl: lines.append(f"domain:{row['host']}")
    disavow_txt=os.path.join(outdir,f"disavow_{CDOM}.txt")
    open(disavow_txt,"w",encoding="utf-8").write("\n".join(lines)+"\n")

    # ================= EXCLUDED ANCHORS LIST =================
    ex=[]
    ex.append(f"# Anchors deliberately EXCLUDED from the {CDOM} disavow file — with the reason for each")
    ex.append("")
    ex.append(f"Generated {datetime.date.today().isoformat()}. This lists every anchor group held back from the one-time")
    ex.append(f"disavow file and why. {spam_file:,} domains ARE disavowed (gambling + link-selling/PBN); the groups below are NOT.")
    ex.append("")
    ex.append("## 1. Exact-match 'dumpster rental <Missouri city>' / service anchors — HELD pending supplier confirmation")
    ex.append("")
    ex.append("Reason: these point at deep service/product pages and read like commissioned local-SEO placements. Disavowing")
    ex.append("work the business (or its previous SEO partner) paid for would hurt it, so they are held until the supplier confirms")
    ex.append("who commissioned them. Cities seen: St Louis, Eureka, Arnold, Festus, Valley Park, Mehlville, Fenton, O'Fallon,")
    ex.append("St Charles, St Peters, Florissant, Ballwin, Chesterfield MO. Representative anchors (Semrush domains / backlinks):")
    ex.append("")
    held_anchor_rows=[(a,dn,bn) for (a,dn,bn) in anchors if C.anchor_type(a).startswith("Exact match")]
    for a,dn,bn in sorted(held_anchor_rows,key=lambda x:-x[2])[:60]:
        ex.append(f"  - \"{a}\"  — {dn} domains, {bn} links  [{C.anchor_type(a)}]")
    ex.append(f"  ... {len(held_anchor_rows)} distinct exact-match anchors in total; {held} referring domains held (see workbook tabs 5 & 6).")
    ex.append("")
    ex.append("## 2. Genuine brand / citation anchors — KEPT (never disavow a real citation)")
    ex.append("")
    ex.append("Reason: real business-listing and citation platforms naming the client. Disavowing these would remove legitimate")
    ex.append("local-SEO signal. Examples: bbb.org, yellowpages.com, superpages.com, crunchbase.com, birdeye.com,")
    ex.append("chamberofcommerce.com, citysquares.com, owler.com, givesmart.com (charity fundraiser), dumpster.co.")
    ex.append("Brand/navigational anchors kept: \"scapersllc.com\", \"visit website\", \"website\", \"scapers llc\".")
    ex.append("")
    ex.append("## 3. Nofollow-only spam anchors — NO ACTION (pass no ranking value)")
    ex.append("")
    ex.append(f"Reason: {spam_nf} domains are spam but every link from them is nofollow, so Google passes no value; disavowing is unnecessary.")
    ex.append("")
    ex.append("## 4. Low-volume, no-visible-anchor domains — KEPT out of the file")
    ex.append("")
    keepn=dc.get("Keep",0)
    ex.append(f"Reason: {keepn} domains (mostly 1-2 links, empty/image anchors) show no visible per-domain spam evidence. Per the")
    ex.append("playbook, nothing is disavowed on a name, TLD, score or shared IP alone, so these are left out. The client can extend")
    ex.append("the file later if Search Console shows them causing harm.")
    open(os.path.join(outdir,"excluded_anchors.md"),"w",encoding="utf-8").write("\n".join(ex)+"\n")

    # ================= SUMMARY STATS =================
    s=[]
    s.append("SCAPERS LLC — backlink audit summary (counts)")
    s.append(f"  Referring domains: {total_d}   Backlinks: {total_b}")
    for st in stat_order:
        if dc.get(st,0): s.append(f"  {st}: {dc[st]} domains / {bc[st]} links")
    s.append(f"  Disavow file lines (deduped): {len(disavow)}")
    s.append(f"  Held (excluded, pending supplier): {held}")
    for grp in order:
        if grp in groups: s.append(f"    disavow group '{grp}': {len(groups[grp])}")
    open(os.path.join(outdir,"audit_summary.txt"),"w",encoding="utf-8").write("\n".join(s)+"\n")
    print("\n".join(s))
    print("\nWritten:", xlsx, "|", disavow_txt, "| excluded_anchors.md | audit_summary.txt")

def dedupe_disavow(spam):
    hosts={r["host"] for r in spam}
    out=[]
    for r in spam:
        h=r["host"]; root=C.host_root(h)
        if h!=root and root in hosts: continue
        out.append(r)
    seen=set(); ded=[]
    for r in sorted(out,key=lambda x:-x["backlinks"]):
        if r["host"] in seen: continue
        seen.add(r["host"]); ded.append(r)
    return ded

if __name__=="__main__": main()
