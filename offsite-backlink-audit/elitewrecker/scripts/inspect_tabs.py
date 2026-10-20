import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])

# Full Summary values rows 1-28
ws = wb["Summary"]
print("##### SUMMARY FULL VALUES #####")
for r in range(1, ws.max_row+1):
    vals = []
    for c in range(1,6):
        v = ws.cell(row=r,column=c).value
        if v is not None: vals.append(f"{chr(64+c)}{r}={v!r}")
    if vals: print(" | ".join(vals))

for name in ["Anchor Analysis","Exact Match Anchors","Pending Disavow Lines","Disavow List"]:
    ws = wb[name]
    print("\n##### ", name, " dims", ws.dimensions, "freeze", ws.freeze_panes, "filter", ws.auto_filter.ref)
    widths = {col:round(d.width,2) for col,d in ws.column_dimensions.items() if d.width}
    print("widths:", widths)
    # header
    hdr = [ws.cell(row=1,column=c).value for c in range(1,ws.max_column+1)]
    print("HEADER:", hdr)
    # 3 sample rows
    for r in range(2, min(ws.max_row,4)+1):
        row=[ws.cell(row=r,column=c).value for c in range(1,ws.max_column+1)]
        print(f"row{r}:", [str(x)[:40] for x in row])
