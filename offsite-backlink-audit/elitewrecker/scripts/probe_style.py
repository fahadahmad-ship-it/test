import sys
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1])
def cellinfo(ws,coord):
    c=ws[coord]; b=c.border
    def side(s): return s.style if s and s.style else None
    return (f"{coord} val={str(c.value)[:30]!r} font={c.font.name},{c.font.sz},b={c.font.bold},"
            f"color={c.font.color.rgb if c.font.color else None} fill={c.fill.fgColor.rgb if c.fill.patternType else None} "
            f"border[L{side(b.left)},R{side(b.right)},T{side(b.top)},B{side(b.bottom)}] "
            f"align={c.alignment.horizontal},{c.alignment.vertical},wrap={c.alignment.wrap_text} numfmt={c.number_format}")
S=wb["Summary"]
for co in ["A11","A19","A20","B20","D4","E4","A12"]:
    print("SUM",cellinfo(S,co))
RDA=wb["Referring Domain Audit"]
for co in ["A1","A2","C1","C2","B3"]:
    print("RDA",cellinfo(RDA,co))
EMA=wb["Exact Match Anchors"]
for co in ["A1","F1","F2","G2","D2","E2"]:
    print("EMA",cellinfo(EMA,co))
PDL=wb["Pending Disavow Lines"]
for co in ["A1","B2","C2"]:
    print("PDL",cellinfo(PDL,co))
DL=wb["Disavow List"]
for co in ["A1","A2","B2","G2"]:
    print("DL",cellinfo(DL,co))
AA=wb["Anchor Analysis"]
for co in ["A1","A2","C2","D2"]:
    print("AA",cellinfo(AA,co))
# row heights full for each sheet (distinct values)
for name in wb.sheetnames:
    ws=wb[name]
    hs={}
    for r,d in ws.row_dimensions.items():
        if d.height: hs.setdefault(d.height,[]).append(r)
    summ={h:(len(rs)) for h,rs in hs.items()}
    print("HEIGHTS",name, summ, "default", ws.sheet_format.defaultRowHeight)
