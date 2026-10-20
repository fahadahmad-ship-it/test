import sys, os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = sys.argv[1]
outdir = sys.argv[2]
wb = load_workbook(path)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n" + "="*70)
    print("SHEET:", repr(ws.title), "dims", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
    print("freeze:", ws.freeze_panes, "autofilter:", ws.auto_filter.ref, "gridlines:", ws.sheet_view.showGridLines)
    print("merged:", [str(m) for m in ws.merged_cells.ranges][:12])
    # column widths
    widths = {}
    for col, dim in ws.column_dimensions.items():
        if dim.width: widths[col] = round(dim.width,2)
    print("col widths:", widths)
    # row heights (first 12)
    rh = {}
    for r,dim in ws.row_dimensions.items():
        if dim.height: rh[r]=dim.height
    print("row heights (sample):", {k:rh[k] for k in sorted(rh)[:12]})
    # images
    imgs = getattr(ws, "_images", [])
    print("images:", len(imgs))
    for i,im in enumerate(imgs):
        try:
            anc = im.anchor
            print("  image", i, "anchor:", anc._from.col, anc._from.colOff, anc._from.row, anc._from.rowOff, "ext", getattr(im,'width',None), getattr(im,'height',None))
        except Exception as e:
            print("  image", i, "anchor-info-err", e)
        data = im._data() if hasattr(im,'_data') else None
        if data:
            fn = os.path.join(outdir, f"{ws.title.replace(' ','_')}_img{i}.png")
            open(fn,'wb').write(data); print("   saved", fn, len(data),"bytes")
    # first 4 rows: values + style of up to 18 cols
    for r in range(1, min(ws.max_row,4)+1):
        for c in range(1, min(ws.max_column,18)+1):
            cell = ws.cell(row=r, column=c)
            if cell.value is None and c>1 and r>1: continue
            f = cell.font; fill = cell.fill; al = cell.alignment
            fg = fill.fgColor.rgb if fill and fill.patternType else None
            print(f"  {cell.coordinate}: val={repr(cell.value)[:50]} font={f.name},{f.sz},b={f.bold},color={f.color.rgb if f.color else None} fill={fg} align={al.horizontal},{al.vertical},wrap={al.wrap_text} numfmt={cell.number_format}")
