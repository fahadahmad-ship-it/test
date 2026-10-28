import sys
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=False)
print("SHEETS (in order):")
for i, ws in enumerate(wb.sheetnames):
    s = wb[ws]
    print(f"  [{i}] {ws!r}  dims={s.dimensions}  max_row={s.max_row} max_col={s.max_column}  gridlines={s.sheet_view.showGridLines} freeze={s.freeze_panes}")
print()
for ws in wb.sheetnames:
    s = wb[ws]
    print("=" * 90)
    print(f"TAB: {ws!r}  (max_row={s.max_row}, max_col={s.max_column})")
    print("-" * 90)
    nrows = min(s.max_row, 8)
    for r in range(1, nrows + 1):
        cells = []
        for c in range(1, min(s.max_column, 20) + 1):
            cell = s.cell(row=r, column=c)
            v = cell.value
            if v is None:
                continue
            vs = str(v).replace("\n", "\\n")
            if len(vs) > 40:
                vs = vs[:40] + "…"
            cells.append(f"{cell.coordinate}={vs!r}")
        if cells:
            print(f"  R{r}: " + " | ".join(cells))
    print()
