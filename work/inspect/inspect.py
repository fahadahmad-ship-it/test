import sys, openpyxl
from openpyxl.utils import get_column_letter
p = sys.argv[1]
wb = openpyxl.load_workbook(p, data_only=False)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n" + "="*70)
    print(f"TAB: {ws.title!r}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("  freeze:", ws.freeze_panes, " autofilter:", ws.auto_filter.ref, " gridlines:", ws.sheet_view.showGridLines)
    # header row(s): print first 3 rows
    for r in range(1, min(ws.max_row,3)+1):
        vals = []
        for c in range(1, ws.max_column+1):
            v = ws.cell(r,c).value
            if v is not None:
                vals.append(f"{get_column_letter(c)}{r}={v!r}")
        print(f"  row{r}:", " | ".join(vals)[:500])
