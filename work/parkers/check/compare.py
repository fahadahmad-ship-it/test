"""compare.py — structural/style comparison of the built workbook vs the K&K template.
Compares layout and styling (not data values, which differ by client): sheet order,
header styles, column widths, freeze panes, autofilter presence, gridlines, data-row
style, Summary title style, logo presence."""
import sys, openpyxl
from openpyxl.utils import get_column_letter

TMPL = "/home/user/test/work/parkers/ref/KK_template.xlsx"
OUT  = sys.argv[1] if len(sys.argv)>1 else "/home/user/test/work/parkers/out/parkersair_backlink_audit.xlsx"
t = openpyxl.load_workbook(TMPL); o = openpyxl.load_workbook(OUT)
issues=[]
def bs(s): return s.style if s else None
def fkey(c):
    b=c.border
    return (c.font.name, c.font.size, c.font.bold,
            (c.font.color.rgb if c.font.color else None),
            (c.fill.fgColor.rgb if c.fill and c.fill.patternType else None),
            c.alignment.horizontal, c.alignment.vertical, c.alignment.wrap_text,
            (bs(b.left), bs(b.right), bs(b.top), bs(b.bottom)))

if t.sheetnames != o.sheetnames:
    issues.append(f"SHEET ORDER differs: {t.sheetnames} vs {o.sheetnames}")

for name in t.sheetnames:
    if name not in o.sheetnames: continue
    ts, os_ = t[name], o[name]
    # gridlines
    if ts.sheet_view.showGridLines != os_.sheet_view.showGridLines:
        issues.append(f"[{name}] gridlines {ts.sheet_view.showGridLines} vs {os_.sheet_view.showGridLines}")
    # freeze
    if ts.freeze_panes != os_.freeze_panes:
        issues.append(f"[{name}] freeze {ts.freeze_panes} vs {os_.freeze_panes}")
    # autofilter presence
    if bool(ts.auto_filter.ref) != bool(os_.auto_filter.ref):
        issues.append(f"[{name}] autofilter {ts.auto_filter.ref} vs {os_.auto_filter.ref}")
    ncol = ts.max_column
    # header styles + values
    for c in range(1, ncol+1):
        tc, oc = ts.cell(1,c), os_.cell(1,c)
        if name!="Summary":
            if tc.value != oc.value:
                issues.append(f"[{name}] header {get_column_letter(c)}1 value {tc.value!r} vs {oc.value!r}")
            if fkey(tc)!=fkey(oc):
                issues.append(f"[{name}] header {get_column_letter(c)}1 STYLE differs")
    # column widths
    for c in range(1, ncol+1):
        cl = get_column_letter(c)
        tw = ts.column_dimensions[cl].width; ow = os_.column_dimensions[cl].width
        if tw and ow and abs(tw-ow)>0.5:
            issues.append(f"[{name}] width {cl} {tw:.1f} vs {ow:.1f}")
    # data-row style (row2) for data tabs
    if name!="Summary" and ts.max_row>=2 and os_.max_row>=2:
        for c in range(1, ncol+1):
            if fkey(ts.cell(2,c))!=fkey(os_.cell(2,c)):
                issues.append(f"[{name}] data-row style {get_column_letter(c)}2 differs")
                break
    # header row height
    if ts.row_dimensions[1].height and os_.row_dimensions[1].height:
        if abs(ts.row_dimensions[1].height-os_.row_dimensions[1].height)>0.3:
            issues.append(f"[{name}] header height {ts.row_dimensions[1].height} vs {os_.row_dimensions[1].height}")

# Summary title + logo
if fkey(t["Summary"]["B1"])!=fkey(o["Summary"]["B1"]):
    issues.append("[Summary] B1 title style differs")
if len(getattr(t["Summary"],'_images',[]))!=len(getattr(o["Summary"],'_images',[])):
    issues.append(f"[Summary] logo count {len(t['Summary']._images)} vs {len(o['Summary']._images)}")

print(f"TEMPLATE COMPARE: {len(issues)} structural/style differences")
for i in issues[:60]: print("  -", i)
if not issues: print("  PASS: zero structural differences.")
