import openpyxl, os, sys, re
from collections import Counter
BASE="/home/user/test/work/parkers"
wb=openpyxl.load_workbook(os.path.join(BASE,"out","parkersair_backlink_audit.xlsx"))
issues=[]
# --- tab row counts ---
def nrows(ws): return ws.max_row-1
rda=wb["Referring Domain Audit"]; allb=wb["All Backlinks"]; anc=wb["Anchor Analysis"]
em=wb["Exact Match Anchors"]; pdl=wb["Pending Disavow Lines"]; dl=wb["Disavow List"]; s=wb["Summary"]
print("ROW COUNTS: RDA=%d AllBL=%d Anchor=%d ExactMatch=%d Pending=%d Disavow=%d"%(
    nrows(rda),nrows(allb),nrows(anc),nrows(em),nrows(pdl),nrows(dl)))
# --- Summary status table reconcile ---
print("\nSUMMARY status table:")
for r in range(4,11):
    print("  ", [s.cell(r,c).value for c in range(1,5)])
# --- status distribution in RDA ---
st=Counter(rda.cell(r,2).value for r in range(2,rda.max_row+1))
print("\nRDA status counts:",dict(st))
# --- metrics populated? Traffic col F (6) in RDA, Disavow List Traffic(5)/Keywords(6) ---
rda_traffic_nonzero=sum(1 for r in range(2,rda.max_row+1) if (rda.cell(r,6).value or 0))
print("RDA rows with Traffic>0:",rda_traffic_nonzero)
dl_kw_present=sum(1 for r in range(2,dl.max_row+1) if dl.cell(r,6).value not in (None,""))
print("Disavow rows with Keywords cell populated:",dl_kw_present,"/",nrows(dl))
# held traffic distribution (join RDA)
held_traffic=[(rda.cell(r,6).value or 0) for r in range(2,rda.max_row+1) if str(rda.cell(r,2).value).startswith("Held")]
print("HELD domains: %d, with 0 traffic: %d, with >0: %d"%(len(held_traffic),sum(1 for t in held_traffic if not t),sum(1 for t in held_traffic if t)))
# --- First Seen populated (no blanks) ---
blank_fs=sum(1 for r in range(2,rda.max_row+1) if rda.cell(r,11).value in (None,""))
print("RDA First Seen blanks:",blank_fs)
# --- disavow line col Q only for spam ---
badQ=0
for r in range(2,rda.max_row+1):
    stt=rda.cell(r,2).value; q=rda.cell(r,17).value
    if stt=="Spam (Disavow)" and not (q and str(q).startswith("domain:")) and q!="redundant": badQ+=1
    if stt in ("Keep","Own Site","Held (Pending Partner)","Held (Nofollow, No Action)","Spam (Nofollow, No Action)") and q: badQ+=1
print("RDA Disavow-Line column anomalies:",badQ)

# --- Disavow file audit ---
f=open(os.path.join(BASE,"out","disavow_parkersair.txt")).read().splitlines()
dom_lines=[l for l in f if l.startswith("domain:")]
comments=[l for l in f if l.startswith("#")]
print("\nDISAVOW FILE: %d domain lines, %d comment lines"%(len(dom_lines),len(comments)))
# all lowercase, valid, no dup
lc=all(l==l.lower() for l in dom_lines)
dups=[d for d,c in Counter(dom_lines).items() if c>1]
valid=all(re.fullmatch(r"domain:[a-z0-9.\-]+", l) for l in dom_lines)
print("  all lowercase:",lc," valid format:",valid," duplicates:",len(dups))
# cross-check: no Held/Keep/Own domain in file
sys.path.insert(0,os.path.join(BASE,"scripts"))
import load as L, classify as C
from collections import defaultdict
bl,_=L.load_backlinks(os.path.join(BASE,"data","fetch")); rd,_=L.load_refdomains(os.path.join(BASE,"data","rd"))
rm={r["domain"]:r for r in rd}; by=defaultdict(list)
for b in bl:
    if b["refdomain"]: by[b["refdomain"]].append(b)
def reg(h):
    p=h.split(".");return h if len(p)<=2 else ".".join(p[-2:])
status={}
for d,links in by.items():
    stt,_,_,_=C.classify_domain(d,links,rm.get(d,rm.get(reg(d),{}))); status[d]=stt
filedoms=set(l.split(":",1)[1] for l in dom_lines)
leak=[d for d in filedoms if status.get(d,"")!="Spam (Disavow)"]
print("  NON-spam domains leaked into file:",len(leak), leak[:8])
missing=[d for d,stt in status.items() if stt=="Spam (Disavow)" and d not in filedoms]
# missing should only be redundant children
redun=set(d for d,stt in status.items() if stt=="Spam (Disavow)" and reg(d)!=d and reg(d) in {x for x,y in status.items() if y=="Spam (Disavow)"})
print("  Spam(Disavow) not in file:",len(missing)," of which redundant children:",len([d for d in missing if d in redun]))
print("\nISSUES:",issues if issues else "none")
