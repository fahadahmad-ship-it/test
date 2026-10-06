import csv, re, subprocess, sys
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO='/home/user/test/'
def host(u):
    try: return urlparse(u).netloc.lower().replace('www.','')
    except: return ''

# --- disavow domains + section ---
ent={}; sec=None
for l in open(REPO+'backlink-audit/disavow-v2-ahrefs.txt',encoding='utf-8'):
    m=re.match(r'# SECTION (\d+)',l)
    if m: sec=int(m.group(1))
    if l.startswith('domain:'): ent[l[7:].split('#')[0].strip().lower()]=sec

SECNAME={1:'S1 Campaign 148096',2:'S2 PBN service pages',3:'S3 High-DA vendor',
         4:'S4 Domain marketplace',6:'S6 Semrush-only',7:'S7 Network D (compromised)',
         8:'S8 Vendor path'}
ACTION={1:'SUBMIT',2:'SUBMIT',3:'SUBMIT',4:'SUBMIT',6:'TIER - see col N',7:'SUBMIT',8:'SUBMIT'}

rows=[]
# --- Ahrefs link rows ---
dr={}
for r in csv.DictReader(open(REPO+'backlink-audit/data/ahrefs/ahrefs-refdomains.tsv',encoding='utf-8-sig',newline=''),delimiter='\t'):
    dr[(r.get('Domain') or '').lower()]=r.get('DR')
seen=set()
for r in csv.DictReader(open(REPO+'backlink-audit/data/ahrefs/ahrefs-backlinks.tsv',encoding='utf-8-sig',newline=''),delimiter='\t'):
    h=host(r.get('Referring page URL',''))
    if h not in ent: continue
    key=(r.get('Referring page URL'),r.get('Anchor'),r.get('Target URL'))
    if key in seen: continue
    seen.add(key)
    s=ent[h]
    rows.append([s,SECNAME.get(s,'?'),h,
        r.get('Referring page URL',''), r.get('Anchor',''), r.get('Target URL',''),
        'nofollow' if str(r.get('Nofollow')).lower()=='true' else 'DOFOLLOW',
        dr.get(h,''), 'yes' if str(r.get('Is spam')).lower()=='true' else 'no',
        (r.get('Referring page title') or '')[:120],
        r.get('First seen',''), r.get('Last seen',''), 'Ahrefs', ACTION.get(s,'')])

# --- Semrush link rows ---
CL=re.compile(r'(?<![a-z0-9.-])ngwindows\.com')
raw=subprocess.run(['git','-C',REPO,'show','6f5f407:backlink-audit/work/links_raw.tsv'],
                   capture_output=True,text=True).stdout.splitlines()
for l in raw:
    p=l.split('\t')
    if len(p)<3: continue
    h=host(p[0])
    if h not in ent: continue
    if p[2].strip().lower()!='false': continue
    if not CL.search(p[1].lower()): continue
    key=(p[0],p[1],'')
    if key in seen: continue
    seen.add(key)
    s=ent[h]
    tier=''
    if s==6:
        tier='SUBMIT - campaign 148096 in path' if '148096' in p[0] else 'HOLD to 25 Oct - anchor only'
    rows.append([s,SECNAME.get(s,'?'),h,p[0],p[1],'(not in Semrush export)','DOFOLLOW',
                 dr.get(h,''),'n/a','', '','','Semrush', tier or ACTION.get(s,'')])

rows.sort(key=lambda r:(r[0], r[2], r[3]))

wb=Workbook(); ws=wb.active; ws.title='Links to disavow'
HDR=['Sec','Section','Referring domain','BACKLINK URL (the spam page)','ANCHOR TEXT',
     'TARGET URL (page on client site)','Follow','DR','Ahrefs spam?','Referring page title',
     'First seen','Last seen','Source','Action']
ws.append(HDR)
for r in rows: ws.append(r)

F='Arial'
hf=Font(name=F,bold=True,color='FFFFFF',size=10); hfill=PatternFill('solid',fgColor='1F3864')
thin=Side(style='thin',color='D9D9D9')
for c in ws[1]:
    c.font=hf; c.fill=hfill; c.alignment=Alignment(vertical='center',wrap_text=True)
ws.row_dimensions[1].height=32
for i in range(2,len(rows)+2):
    for j in range(1,len(HDR)+1):
        c=ws.cell(row=i,column=j); c.font=Font(name=F,size=9)
        c.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        c.alignment=Alignment(vertical='top',wrap_text=(j in (4,5,6,10)))
    if ws.cell(row=i,column=7).value=='DOFOLLOW':
        ws.cell(row=i,column=7).fill=PatternFill('solid',fgColor='FFC7CE')
    a=str(ws.cell(row=i,column=14).value or '')
    if a.startswith('HOLD'): ws.cell(row=i,column=14).fill=PatternFill('solid',fgColor='FFE699')
    elif a.startswith('SUBMIT'): ws.cell(row=i,column=14).fill=PatternFill('solid',fgColor='C6EFCE')
W={'A':5,'B':24,'C':30,'D':62,'E':52,'F':40,'G':11,'H':7,'I':11,'J':44,'K':12,'L':12,'M':9,'N':30}
for k,v in W.items(): ws.column_dimensions[k].width=v
ws.freeze_panes='D2'
ws.auto_filter.ref=f"A1:{get_column_letter(len(HDR))}{len(rows)+1}"

# --- domain list sheet (what you upload) ---
d=wb.create_sheet('Disavow file (domains)')
d.append(['Line to paste into Google','Section','Action'])
for c in d[1]: c.font=hf; c.fill=hfill
for dom in sorted(ent, key=lambda x:(ent[x],x)):
    s=ent[dom]
    d.append([f'domain:{dom}',SECNAME.get(s,''),ACTION.get(s,'')])
for i in range(2,len(ent)+2):
    for j in (1,2,3): d.cell(row=i,column=j).font=Font(name=F,size=10)
d.column_dimensions['A'].width=44; d.column_dimensions['B'].width=26; d.column_dimensions['C'].width=18
d.freeze_panes='A2'; d.auto_filter.ref=f"A1:C{len(ent)+1}"

print(f"{out_rows if False else len(rows)} link rows, {len(ent)} domains")
wb.save(sys.argv[1])
