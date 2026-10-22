"""
Build the workbook + disavow file(s) for the aaronshometech.com audit, matching
the signed K&K template (ref/KandK_template.xlsx) cell-for-cell: column-A start,
header-row-only borders (data cells unbordered), 21pt client-name title, the two
Summary tables (status + anchor-type profile) with the Harbinger logo, and the
exact per-tab column widths, row heights and number formats.

Run: python3 -I scripts/build.py data/classified.json scripts/config.py \
        data/anchors_0.csv out/Aarons_Home_Tech_Backlink_Audit.xlsx out/
"""
import sys, os, json, glob, html, datetime, importlib.util, re
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

ACCENT = "FF0C343D"; WHITE = "FFFFFFFF"; BLACK = "FF000000"; FONT = "Geist"
THIN = Side(style="thin", color="FF000000")
HDR_BORDER = Border(left=THIN, right=THIN, top=THIN)          # header row: L/R/T, no bottom
FULL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)  # summary tables
NO_BORDER = Border()
WHITE_FILL = PatternFill("solid", fgColor=WHITE)
ACCENT_FILL = PatternFill("solid", fgColor=ACCENT)
HDR_FONT = Font(name=FONT, size=9, bold=True, color=WHITE)
HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(name=FONT, size=9, color=BLACK)
DATA_ALIGN = Alignment(vertical="center")
CTR_ALIGN = Alignment(horizontal="center", vertical="center")
INPUT_FONT = Font(name=FONT, size=9, color=WHITE)

# columns (1-indexed) the template centre-aligns per tab; text columns stay default
CENTER = {
 "Referring Domain Audit": {3,4,5,6,7,9,10,11,15,16},
 "All Backlinks": {6},
 "Anchor Analysis": {3,4,5,6},
 "Exact Match Anchors": {3,4},
 "Pending Disavow Lines": set(),
 "Disavow List": {2,3,4,5,6},
}

WIDTHS = {
 "Referring Domain Audit": [36,24.71,9.57,13,9.43,9.71,12.43,21.29,14,8.71,12,45,8.71,40,14,15,14],
 "All Backlinks": [25.43,27.14,14.71,55,32,10,55,53.14,30.57],
 "Anchor Analysis": [115.86,37.71,14,8.71,None,20.29],
 "Exact Match Anchors": [25.43,59,53.14,16,24.57,15.43,15],
 "Pending Disavow Lines": [32,44,48],
 "Disavow List": [44,14,15.57,14,11.14,14.86,98.86],
}
NUMFMT = {
 "Referring Domain Audit": {4:"#,##0",5:"#,##0",6:"#,##0",11:"d mmm yyyy"},
 "Anchor Analysis": {3:"#,##0",4:"0.00%"},
 "Disavow List": {3:"#,##0",4:"#,##0",5:"#,##0",6:"#,##0"},
}

def load_cfg(p):
    spec = importlib.util.spec_from_file_location("config", p); m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m
def load_at():
    spec = importlib.util.spec_from_file_location("anchortype", "scripts/anchortype.py"); m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m

def ts_to_date(ts):
    try:
        ts = int(ts)
        if ts <= 0: return None
        return datetime.datetime.fromtimestamp(ts, datetime.timezone.utc).date()
    except (ValueError, TypeError, OSError):
        return None

def header(ws, headers, left_cols=(), nowrap_cols=()):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = HDR_FONT; c.fill = ACCENT_FILL; c.border = HDR_BORDER
        horiz = "left" if j in left_cols else "center"
        wrap = j not in nowrap_cols
        c.alignment = Alignment(horizontal=horiz, vertical="center", wrap_text=wrap)
    ws.row_dimensions[1].height = 31.5

LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
def data_row(ws, values, row, height, input_cols=(), numfmt=None, center_cols=frozenset(), left_cols=frozenset()):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        if (j - 1) in input_cols:
            c.font = INPUT_FONT; c.fill = ACCENT_FILL
        else:
            c.font = DATA_FONT; c.fill = WHITE_FILL
        c.alignment = CTR_ALIGN if j in center_cols else (LEFT_ALIGN if j in left_cols else DATA_ALIGN)
        c.border = NO_BORDER
        if numfmt and j in numfmt:
            c.number_format = numfmt[j]
    ws.row_dimensions[row].height = height

def finalize(ws, title, nrows, ncols):
    last = get_column_letter(ncols)
    ws.auto_filter.ref = f"A1:{last}{1+nrows}"
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = True
    for i, w in enumerate(WIDTHS.get(title, []), start=1):
        if w:
            ws.column_dimensions[get_column_letter(i)].width = w

def load_metrics(base):
    cmp_map, rank_map = {}, {}
    for p in glob.glob(os.path.join(base, "metrics", "*cmp*.csv")):
        for line in open(p, encoding="utf-8", errors="replace"):
            line = html.unescape(line.strip())
            if not line or line.startswith("target;"): continue
            x = line.split(";")
            if len(x) < 4: continue
            try: cmp_map[x[0].strip().lower()] = {"as": int(x[1]), "backlinks": int(x[2]), "rds": int(x[3])}
            except ValueError: pass
    for p in glob.glob(os.path.join(base, "metrics", "*rank*.csv")):
        for line in open(p, encoding="utf-8", errors="replace"):
            line = html.unescape(line.strip())
            if not line or line.lower().startswith("domain;"): continue
            x = line.split(";")
            if len(x) < 4: continue
            try: rank_map[x[0].strip().lower()] = {"organic_keywords": int(x[2]), "organic_traffic": int(x[3])}
            except (ValueError, IndexError): pass
    return cmp_map, rank_map

def build_disavow_set(domains):
    spam = {d: r for d, r in domains.items() if r["status"] == "Spam (Disavow)" and r["has_dofollow"]}
    keys = set(spam)
    redundant = set()
    for d in spam:
        labs = d.split(".")
        for i in range(1, len(labs) - 1):
            if ".".join(labs[i:]) in keys:
                redundant.add(d); break
    return {d: r for d, r in spam.items() if d not in redundant}, redundant

HOME_RE = re.compile(r"^https?://(www\.)?aaronshometech\.com/?$", re.I)

def main():
    classified, cfg_p, anchors_csv, out_xlsx = sys.argv[1:5]
    out_dir = sys.argv[5] if len(sys.argv) > 5 else "out"
    cfg = load_cfg(cfg_p); AT = load_at(); base = "data"
    data = json.load(open(classified, encoding="utf-8"))["domains"]
    cmp_map, rank_map = load_metrics(base)
    doms = list(data.values())
    def AS(r): return cmp_map.get(r["domain"], {}).get("as", r.get("domain_ascore", 0))
    def own_bl(r): return cmp_map.get(r["domain"], {}).get("backlinks", "")
    def own_rd(r): return cmp_map.get(r["domain"], {}).get("rds", "")
    # organic traffic/keywords: real value where pulled; 0 for no-authority spam not
    # individually queried (AS<3, validated: 0/60 sampled had any organic traffic).
    def traf(r): return rank_map.get(r["domain"], {}).get("organic_traffic", 0)
    def okw(r): return rank_map.get(r["domain"], {}).get("organic_keywords", 0)
    disavow, redundant = build_disavow_set(data)
    disset = set(disavow)

    wb = openpyxl.Workbook()
    rank_sort = {"Spam (Disavow)":0,"Held (Pending Partner)":1,"Held (Nofollow, No Action)":2,"Spam (Nofollow, No Action)":3,"Keep":4,"Own Site":5}
    doms_sorted = sorted(doms, key=lambda r: (rank_sort.get(r["status"],9), -AS(r), r["domain"]))

    # ===== SUMMARY =====
    ws = wb.active; ws.title = "Summary"; ws.sheet_view.showGridLines = False
    for i, w in enumerate([38.86,16.29,17.57,18.71,120.0], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    if os.path.exists("data/harbinger_logo.png"):
        img = XLImage("data/harbinger_logo.png"); img.width = 170; img.height = 50
        ws.add_image(img, "A1")
    ws.merge_cells("B1:E2")
    t = ws.cell(row=1, column=2, value=cfg.CLIENT_NAME)
    t.font = Font(name=FONT, size=21, bold=True, color=ACCENT)
    t.alignment = Alignment(horizontal="center", vertical="center")

    STATUS_DESC = {
     "Spam (Disavow)":"Confirmed spam with at least one dofollow link. Included in the disavow file.",
     "Held (Pending Partner)":"Exact match commercial anchor, or a link aimed at a product, capability or industry page. May have been commissioned. Not in the disavow file until the previous partner confirms.",
     "Held (Nofollow, No Action)":"Same, but every link is nofollow so it passes no ranking value. No disavow needed either way.",
     "Spam (Nofollow, No Action)":"Spam, but every link is nofollow and passes no ranking value, so no disavow is needed.",
     "Keep":"Genuine or harmless link. Keep.",
     "Own Site":"Client owned domain that redirects into the site.",
    }
    order = ["Spam (Disavow)","Held (Pending Partner)","Held (Nofollow, No Action)","Spam (Nofollow, No Action)","Keep","Own Site"]
    counts = Counter(r["status"] for r in doms)
    bl_by_status = defaultdict(int)
    for r in doms: bl_by_status[r["status"]] += int(r.get("backlinks_num", 0))
    total_dom = len(doms); total_bl = sum(bl_by_status.values())

    def sumcell(r, c, val, bold=False, fill=True, fmt=None, wrap=False, border=True):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = Font(name=FONT, size=9, bold=bold, color=BLACK)
        if fill: cell.fill = WHITE_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=wrap)
        if border: cell.border = FULL_BORDER
        if fmt: cell.number_format = fmt
        return cell

    # status table header (row 3)
    for j, h in enumerate(["Status","Domains","Backlinks","Share of Domains","What it means"], start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color=WHITE); c.fill = ACCENT_FILL
        c.alignment = HDR_ALIGN; c.border = FULL_BORDER
    r = 3
    for st in order:
        r += 1
        n = counts.get(st, 0)
        sumcell(r,1,st); sumcell(r,2,n,fmt="#,##0"); sumcell(r,3,bl_by_status.get(st,0),fmt="#,##0")
        sumcell(r,4,(n/total_dom if total_dom else 0),fmt="0.0%"); sumcell(r,5,STATUS_DESC[st],wrap=True)
    r += 1
    sumcell(r,1,"Total",bold=True); sumcell(r,2,total_dom,bold=True,fmt="#,##0")
    sumcell(r,3,total_bl,bold=True,fmt="#,##0"); sumcell(r,4,1.0,bold=True,fmt="0.0%"); sumcell(r,5,"",)
    # key findings
    r += 1
    kf = ws.cell(row=r, column=1, value="Key findings"); kf.font = Font(name=FONT, size=10, bold=True, color=BLACK)
    findings = compute_findings(doms, disavow, cfg, counts, total_dom, bl_by_status)
    for f in findings:
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(row=r, column=1, value=f); c.font = Font(name=FONT, size=9, color=BLACK)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    # anchor-type profile
    r += 1
    sec = ws.cell(row=r, column=1, value="Anchor text profile (all backlinks)")
    sec.font = Font(name=FONT, size=11, bold=True, color=ACCENT)
    r += 1
    for j, h in enumerate(["Anchor type","Backlinks","Share of Backlinks","Referring Domains"], start=1):
        c = ws.cell(row=r, column=j, value=h); c.font = Font(name=FONT, size=9, bold=True, color=WHITE)
        c.fill = ACCENT_FILL; c.alignment = HDR_ALIGN; c.border = FULL_BORDER
    anchors = read_anchor_csv(anchors_csv)
    at_bl = defaultdict(int); at_rd = defaultdict(int)
    tot_abl = sum(a["backlinks_num"] for a in anchors) or 1
    for a in anchors:
        ty = AT.anchor_type(a["anchor"], cfg)
        at_bl[ty] += a["backlinks_num"]; at_rd[ty] += a["domains_num"]
    for ty in AT.TYPES:
        if at_bl.get(ty, 0) == 0: continue
        r += 1
        sumcell(r,1,ty); sumcell(r,2,at_bl[ty],fmt="#,##0")
        sumcell(r,3,at_bl[ty]/tot_abl,fmt="0.0%"); sumcell(r,4,at_rd[ty],fmt="#,##0")
    for rr in range(1, r + 1):
        ws.row_dimensions[rr].height = 22.5

    # ===== REFERRING DOMAIN AUDIT =====
    ws = wb.create_sheet("Referring Domain Audit")
    H = ["Domain","Status","AS","Backlinks","RD's","Traffic","In Disavow","Links Point To","Dofollow Links","Nofollow Links","First Seen","Top Anchor Text","Example Linking Page","Linking Page Title","Source","IP","Disavow Line"]
    header(ws, H)
    for i, r in enumerate(doms_sorted, start=2):
        links = r.get("links", [])
        df = sum(1 for l in links if not l["nofollow"]); nf = sum(1 for l in links if l["nofollow"])
        ex = r.get("example", {}) or {}
        indis = "Yes" if r["domain"] in disset else ("Redundant" if r["domain"] in redundant else "No")
        data_row(ws, [r["domain"], r["status"], AS(r), own_bl(r), own_rd(r), traf(r), indis,
                      r.get("points_to",""), df, nf, ts_to_date(r.get("first_seen")),
                      (r.get("top_anchor","") or "")[:200], (ex.get("source_url","") or "")[:300],
                      (ex.get("source_title","") or "")[:300], "Semrush", r.get("ip",""),
                      f"domain:{r['domain']}" if r["domain"] in disset else ""],
                 i, 15.75, numfmt=NUMFMT["Referring Domain Audit"], center_cols=CENTER["Referring Domain Audit"])
    finalize(ws, "Referring Domain Audit", len(doms_sorted), 17)

    # ===== ALL BACKLINKS =====
    ws = wb.create_sheet("All Backlinks")
    header(ws, ["Referring Domain","Status","Tool","Anchor Text","Anchor Type","Link Type","Linking Page URL","Linking Page Title","Target URL"], left_cols={3})
    row = 2
    for r in doms_sorted:
        for l in r.get("links", []):
            data_row(ws, [r["domain"], r["status"], "Semrush", (l.get("anchor","") or "")[:300],
                          AT.anchor_type(l.get("anchor",""), cfg, l.get("target_url","")),
                          "Nofollow" if l["nofollow"] else "Dofollow", (l.get("source_url","") or "")[:300],
                          (l.get("source_title","") or "")[:300], (l.get("target_url","") or "")[:300]], row, 22.5, center_cols=CENTER["All Backlinks"], left_cols={3})
            row += 1
    finalize(ws, "All Backlinks", row - 2, 9)

    # ===== ANCHOR ANALYSIS =====
    ws = wb.create_sheet("Anchor Analysis")
    header(ws, ["Anchor Text","Anchor Type","Backlinks","Share of Backlinks","Referring Domains","Spam Domains Using It"], nowrap_cols={1})
    anchor_spam = defaultdict(set)
    for r in doms:
        if r["status"].startswith("Spam"):
            for l in r.get("links", []):
                anchor_spam[(l.get("anchor","") or "").strip().lower()].add(r["domain"])
    anchors.sort(key=lambda a: -a["backlinks_num"])
    for i, a in enumerate(anchors, start=2):
        su = len(anchor_spam.get(a["anchor"].strip().lower(), set()))
        data_row(ws, [a["anchor"][:300], AT.anchor_type(a["anchor"], cfg), a["backlinks_num"],
                      a["backlinks_num"]/tot_abl, a["domains_num"], su if su else ""], i, 22.5,
                 numfmt=NUMFMT["Anchor Analysis"], center_cols=CENTER["Anchor Analysis"])
    finalize(ws, "Anchor Analysis", len(anchors), 6)

    # ===== EXACT MATCH ANCHORS =====
    ws = wb.create_sheet("Exact Match Anchors")
    header(ws, ["Referring domain","Example backlink (linking page)","Exact-match commercial anchor","Market","Site type (indicative)","Partner confirmed? (Y/N)","Final decision"])
    em = build_exact_match_rows(doms, cfg, AT)
    for i, e in enumerate(em, start=2):
        data_row(ws, e, i, 22.5, input_cols=(5, 6), center_cols=CENTER["Exact Match Anchors"])
    finalize(ws, "Exact Match Anchors", len(em), 7)

    # ===== PENDING DISAVOW LINES =====
    ws = wb.create_sheet("Pending Disavow Lines")
    header(ws, ["Referring domain","Disavow line (append to disavow.txt once confirmed)","Recommendation"])
    held = [r for r in doms_sorted if r["status"] == "Held (Pending Partner)"]
    for i, r in enumerate(held, start=2):
        data_row(ws, [r["domain"], f"domain:{r['domain']}",
                      "Hold — exact-match commercial / deep money-page link; confirm with the previous supplier (partner) whether commissioned. Disavow only if not commissioned."], i, 22.5)
    finalize(ws, "Pending Disavow Lines", len(held), 3)

    # ===== DISAVOW LIST =====
    ws = wb.create_sheet("Disavow List")
    header(ws, ["Disavow Line","AS","Backlinks","RD's","Traffic","Organic Keywords","Top Anchor Text"])
    for i, r in enumerate(sorted(disavow.values(), key=lambda r:(-AS(r), r["domain"])), start=2):
        data_row(ws, [f"domain:{r['domain']}", AS(r), own_bl(r), own_rd(r), traf(r), okw(r),
                      (r.get("top_anchor","") or "")[:200]], i, 22.5, numfmt=NUMFMT["Disavow List"], center_cols=CENTER["Disavow List"])
    finalize(ws, "Disavow List", len(disavow), 7)

    os.makedirs(os.path.dirname(out_xlsx), exist_ok=True)
    wb.save(out_xlsx)
    write_disavow_file(os.path.join(out_dir, "disavow_aaronshometech.com.txt"), disavow, AS, cfg)
    sys.stderr.write(f"[build] {out_xlsx}: domains={len(doms)} disavow={len(disavow)} held={len(held)} anchors={len(anchors)} exact_match={len(em)}\n")


def read_anchor_csv(path):
    out = []
    for i, line in enumerate(open(path, encoding="utf-8", errors="replace")):
        line = html.unescape(line.rstrip("\n").rstrip("\r"))
        if not line: continue
        x = line.split(";")
        if i == 0 and x[0] == "anchor": continue
        if len(x) != 5: continue
        try: out.append({"anchor": x[0], "domains_num": int(x[1]), "backlinks_num": int(x[2])})
        except ValueError: pass
    return out

_CITIES = {"atlanta":"Atlanta, GA","greensboro":"Greensboro, GA","senoia":"Senoia, GA","roswell":"Roswell, GA","decatur":"Decatur, GA","marietta":"Marietta, GA","johns creek":"Johns Creek, GA","smyrna":"Smyrna, GA","sandy springs":"Sandy Springs, GA","vinings":"Vinings, GA","brookhaven":"Brookhaven, GA","peachtree city":"Peachtree City, GA","eatonton":"Eatonton, GA","chattahoochee hills":"Chattahoochee Hills, GA","fayetteville":"Fayetteville, GA","highlands":"Highlands, NC"}

def _market(a, target_url=""):
    a = (a or "").lower()
    for k, v in _CITIES.items():
        if k in a: return v
    # anchor carried no city -> read it from the target money page's slug
    t = (target_url or "").lower().replace("-", " ").replace("/", " ")
    for k, v in _CITIES.items():
        if k in t: return v
    return "GA (metro Atlanta)"

def _site_type(d, t):
    d = d.lower()
    if any(x in d for x in ("directory","listing","listings","local")): return "Business directory"
    if any(x in d for x in ("press","news","einpresswire","newswire","pr")): return "Press release"
    if any(x in d for x in ("blog","article","articles")): return "Blog / article"
    return "Directory / blog (indicative)"

def build_exact_match_rows(doms, cfg, AT):
    rows = []
    for r in doms:
        seen = set()
        for l in r.get("links", []):
            a = (l.get("anchor","") or "").strip()
            if AT.anchor_type(a, cfg, l.get("target_url","")).startswith("Exact match commercial") and a.lower() not in seen:
                seen.add(a.lower())
                rows.append([r["domain"], (l.get("source_url","") or "")[:300], a[:200], _market(a, l.get("target_url","")), _site_type(r["domain"], l.get("source_title","")), "", ""])
    rows.sort(key=lambda x: x[0])
    return rows

def compute_findings(doms, disavow, cfg, counts, total, bl_by_status):
    held = counts.get("Held (Pending Partner)", 0)
    spam_nf = counts.get("Spam (Nofollow, No Action)", 0)
    home = deep = 0
    for r in doms:
        for l in r.get("links", []):
            tu = re.sub(r"^https?://(www\.)?","",(l.get("target_url","") or "").lower())
            if tu.rstrip("/") == "aaronshometech.com": home += 1
            elif tu.startswith("aaronshometech.com/"): deep += 1
    return [
      f"The profile holds {total:,} referring domains. {len(disavow):,} are confirmed spam with at least one dofollow link and go in the disavow file; {spam_nf:,} more are spam but link only via nofollow, so they pass nothing and need no action.",
      "This is an active negative-SEO / link-seller flood: Semrush shows referring domains rising from ~900 a year ago to 1,699 last month to 5,399 now, while Authority Score fell from 28 (2022) to 10. No brief figure was supplied, so the full live profile was audited.",
      "One disavow file is prepared, for the aaronshometech.com Search Console property, and it carries the full union of bad domains. aaronssystems.com and aaronssecurity.com redirect into the profile and are pending the client's confirmation of ownership and upload access.",
      f"Two campaigns show up by target page: the spam flood points almost entirely at the home page ({home:,} analyzed links), while a smaller, older, commissioned-looking set points at deep /services/ and /location/ money pages ({deep:,}).",
      "The dominant anchors are one link vendor's PBN sales copy — 'premium pbn backlinks', 'high quality dofollow backlinks DA 50 PA 40 … buy backlinks online cheap', niche-edit and emoji testimonial text — the same network also selling links to unrelated businesses. ('DA/PA/DR' in those anchors is the vendor's own Moz/Ahrefs-style wording, quoted as evidence; every authority figure in this audit is the Semrush Authority Score shown in the AS column.)",
      f"{held:,} domains are held for the previous supplier (partner) to confirm: these carry exact-match commercial (geo) anchors or links to deep money pages and may be commissioned placements. They stay out of the disavow file until confirmed.",
      "Held does not mean valuable. Many held domains are themselves low-authority and may still be disavowed once the supplier conversation rules out a commission.",
    ][:7]

def write_disavow_file(path, disavow, AS, cfg):
    groups = defaultdict(list)
    for d, r in disavow.items():
        groups[r.get("spam_group") or "pbn"].append(d)
    label = {"pbn":"PBN / link-service sales-copy spam (dofollow)","off_topic":"Off-topic spam (dofollow)","host":"Link-farm / directory host signature (dofollow)"}
    L = [f"# Disavow file for {cfg.MONEY_SITE} (Google Search Console property)",
         f"# Prepared by Harbinger Marketing — {datetime.date.today().isoformat()}",
         f"# {len(disavow)} domains. One domain: line each; a domain: line also covers its subdomains.",
         "# SCOPE: full UNION of bad domains for the money site. Owned domains aaronssystems.com and",
         "#   aaronssecurity.com 301 into this profile; their bad links are folded in here (one file).",
         "#   Per-property files can be produced if the client has Search Console access to them.",
         "# HELD links are NOT in this file. Exact-match commercial and deep-money-page links that may be",
         "#   commissioned are held for the previous SEO supplier (the 'partner' on the sheet; the client's",
         "#   brief calls this role the 'supplier') to confirm first.",
         "# DO NOT UPLOAD before that supplier conversation. Uploading REPLACES any existing disavow file —",
         "#   confirm with the client whether one already exists (none was supplied for this audit).", ""]
    for g in ("pbn","off_topic","host"):
        if not groups.get(g): continue
        L.append(f"# --- {label[g]}: {len(groups[g])} domains ---")
        L += [f"domain:{d}" for d in sorted(groups[g])]; L.append("")
    open(path, "w", encoding="utf-8").write("\n".join(L))
    sys.stderr.write(f"[build] {path} ({len(disavow)} lines)\n")

if __name__ == "__main__":
    main()
