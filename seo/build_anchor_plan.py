#!/usr/bin/env python3
"""Master Lawn — anchor text & target URL plan for authority link acquisition (Task 4)."""
import os, importlib.util
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("bv", os.path.join(BASE, "build_audit_v2.py"))
B = importlib.util.module_from_spec(spec); spec.loader.exec_module(B)

HOME = "https://www.masterlawn.com/"
LOC = "https://www.masterlawn.com/location/lawn-care-services-{}/"

# target URLs with role + current referring domains (from audit/top-pages)
TARGETS = [
    (HOME, "Homepage (primary)", 316, "Highest priority — brand authority"),
    (LOC.format("memphis-tn"), "Location — Memphis TN", 3, "Core metro; weak links"),
    (LOC.format("germantown-tn"), "Location — Germantown TN", 33, "Strong metro term"),
    (LOC.format("collierville-tn"), "Location — Collierville TN", 2, "Affluent metro; grow"),
    (LOC.format("bartlett-tn"), "Location — Bartlett TN", 41, "Core metro"),
    (LOC.format("arlington-tn"), "Location — Arlington TN", 1, "Thin; needs links"),
    (LOC.format("olive-branch-ms"), "Location — Olive Branch MS", 41, "Best-performing local page"),
    (LOC.format("southaven-ms"), "Location — Southaven MS", 1, "Thin; needs links"),
    (LOC.format("huntsville-al"), "Location — Huntsville AL", 2, "PRIORITY — weakest region, locals beat client"),
    ("https://www.masterlawn.com/watering-your-lawn-after-fertilizing-how-long-to-wait-how-long-to-water/",
     "Blog asset (link magnet)", 3, "Ranks 58 kw; earn editorial links"),
    ("https://www.masterlawn.com/how-to-tell-the-difference-between-dead-and-dormant-grass/",
     "Blog asset (link magnet)", 7, "Already attracts links; pitch as resource"),
    ("https://www.masterlawn.com/10-common-spring-lawn-weeds-in-tn-and-north-ms-and-how-to-kill-them/",
     "Blog asset (regional)", 2, "TN/MS weed guide; local relevance"),
]

# recovery-focused anchor mix (rebalance AWAY from spammed exact-match)
RATIO = [
    ("Branded", "45%", "Master Lawn / Master Lawn Inc / Master Lawn (Memphis)",
     "Safest; dominant share to dilute spam over-optimization"),
    ("Naked URL", "22%", "masterlawn.com / www.masterlawn.com / full URL",
     "Natural, penalty-safe"),
    ("Generic / natural", "15%", "visit website / learn more / this Memphis lawn care company",
     "Looks editorial"),
    ("Partial-match / topical", "13%", "lawn care in Germantown, TN / Memphis lawn care experts",
     "Light geo/service context, phrased naturally"),
    ("Exact-match commercial", "5%", "lawn care germantown tn",
     "SPARINGLY — already saturated by the spam; overuse re-triggers over-optimization"),
]

# anchor -> target plan  (target, anchor, type, priority, note)
PLAN = [
    (HOME, "Master Lawn", "Branded", "High", "Primary brand anchor"),
    (HOME, "Master Lawn Inc.", "Branded", "High", "Brand variant"),
    (HOME, "Master Lawn (Memphis, TN)", "Branded", "High", "Brand + geo, natural"),
    (HOME, "masterlawn.com", "Naked URL", "High", "Naked domain"),
    (HOME, "www.masterlawn.com", "Naked URL", "Med", "Naked domain variant"),
    (HOME, "https://www.masterlawn.com/", "Naked URL", "Med", "Full URL"),
    (HOME, "visit website", "Generic", "Med", "Editorial/citation context"),
    (HOME, "learn more", "Generic", "Low", "Editorial context"),
    (HOME, "this Memphis lawn care company", "Partial/natural", "High", "Natural mention in article"),
    (HOME, "professional lawn care and fertilization in Memphis", "Partial/topical", "Med", "Service + geo, natural"),
    (LOC.format("memphis-tn"), "lawn care in Memphis, TN", "Partial/topical", "High", "Natural geo phrasing"),
    (LOC.format("memphis-tn"), "Memphis lawn care services", "Partial/topical", "Med", "Service + geo"),
    (LOC.format("germantown-tn"), "lawn care in Germantown, TN", "Partial/topical", "High", "Natural geo"),
    (LOC.format("germantown-tn"), "Master Lawn — Germantown", "Branded+geo", "Med", "Brand + city"),
    (LOC.format("germantown-tn"), "lawn care germantown tn", "Exact-match", "Low", "Use only on a genuine local directory; sparingly"),
    (LOC.format("collierville-tn"), "Collierville lawn service", "Partial/topical", "High", "Natural geo"),
    (LOC.format("collierville-tn"), "lawn care in Collierville, TN", "Partial/topical", "Med", "Natural geo"),
    (LOC.format("bartlett-tn"), "lawn care serving Bartlett, TN", "Partial/topical", "High", "Natural geo"),
    (LOC.format("bartlett-tn"), "Master Lawn — Bartlett", "Branded+geo", "Med", "Brand + city"),
    (LOC.format("arlington-tn"), "lawn care in Arlington, TN", "Partial/topical", "High", "Thin page — grow"),
    (LOC.format("olive-branch-ms"), "lawn care in Olive Branch, MS", "Partial/topical", "High", "Best local page"),
    (LOC.format("olive-branch-ms"), "Olive Branch lawn service", "Partial/topical", "Med", "Natural geo"),
    (LOC.format("southaven-ms"), "Southaven, MS lawn care", "Partial/topical", "High", "Thin page — grow"),
    (LOC.format("huntsville-al"), "lawn care in Huntsville, AL", "Partial/topical", "TOP", "Weakest region — prioritize"),
    (LOC.format("huntsville-al"), "Huntsville lawn service", "Partial/topical", "TOP", "Weakest region — prioritize"),
    (LOC.format("huntsville-al"), "Master Lawn — Huntsville", "Branded+geo", "High", "Brand + city"),
    ("https://www.masterlawn.com/watering-your-lawn-after-fertilizing-how-long-to-wait-how-long-to-water/",
     "how long to wait to water after fertilizing", "Topical/natural", "High", "Resource-link / guest-post anchor"),
    ("https://www.masterlawn.com/how-to-tell-the-difference-between-dead-and-dormant-grass/",
     "how to tell dead vs dormant grass", "Topical/natural", "High", "Resource-link anchor"),
    ("https://www.masterlawn.com/10-common-spring-lawn-weeds-in-tn-and-north-ms-and-how-to-kill-them/",
     "common spring lawn weeds in Tennessee & North Mississippi", "Topical/natural", "Med", "Regional resource anchor"),
]

def build():
    wb = Workbook()
    F, fill, banner, header_row = B.F, B.fill, B.banner, B.header_row
    NAVY, INK, SOFT, RED, AQUA, BAND, WHITE = B.NAVY, B.INK, B.INK_SOFT, B.RED, B.AQUA, B.BAND, B.WHITE

    # ---- Strategy ----
    ws = wb.active; ws.title = "Strategy"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col, w in zip("BCDE", (22, 10, 52, 46)):
        ws.column_dimensions[col].width = w
    hrow = banner(ws, "ANCHOR & TARGET-URL STRATEGY — Master Lawn (link acquisition)", 5,
                  "Grounded in the audit: the negative-SEO spam over-optimized EXACT-MATCH local anchors, "
                  "so recovery links skew BRANDED / NATURAL. Semrush data · 2026-09-04")
    r = hrow + 1
    ws.cell(row=r, column=2, value="Why this mix").font = F(12, True, NAVY); r += 1
    for txt in [
        "The spam blast used dozens of exact-match local anchors ('lawn care germantown tn', "
        "'mosquito control huntsville al', etc.) — that anchor space is now saturated and toxic-looking.",
        "New authority links must REBALANCE the profile toward branded and naked-URL anchors, with only "
        "light, naturally-phrased geo/service anchors. Piling on more exact-match would re-trigger "
        "over-optimization and undo the disavow's benefit.",
        "Every target below is a real, indexable client page (homepage, the 8 location pages, and the "
        "blog assets that already earn links).",
    ]:
        c = ws.cell(row=r, column=2, value="• " + txt); c.font = F(10, False, INK)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5); ws.row_dimensions[r].height = 42; r += 1
    r += 1
    cols = ["Anchor type", "Target share", "Examples", "Rationale"]
    for i, c in enumerate(cols, 2):
        cell = ws.cell(row=r, column=i, value=c); cell.font = F(10, True, WHITE); cell.fill = fill(B.NAVY2)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    for typ, share, ex, why in RATIO:
        ws.cell(row=r, column=2, value=typ).font = F(10, True)
        ws.cell(row=r, column=3, value=share).font = F(10, True, RED if "Exact" in typ else NAVY)
        ws.cell(row=r, column=4, value=ex).font = F(10)
        ws.cell(row=r, column=5, value=why).font = F(9, False, SOFT)
        for cc in range(2, 6):
            ws.cell(row=r, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30; r += 1

    # ---- Target URLs ----
    ws = wb.create_sheet("Target URLs")
    cols = ["#", "Target URL", "Page role", "Ref domains (now)", "Priority note"]
    hrow = banner(ws, "TARGET URLs — where new links should point", len(cols))
    header_row(ws, cols, hrow, [5, 60, 26, 14, 44])
    for i, (u, role, rd, note) in enumerate(TARGETS, 1):
        ws.append([i, u, role, rd, note])
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(cols))}{ws.max_row}"

    # ---- Anchor–Target Plan ----
    ws = wb.create_sheet("Anchor-Target Plan")
    cols = ["#", "Target URL", "Suggested anchor text", "Anchor type", "Priority", "Use-case / notes"]
    hrow = banner(ws, "ANCHOR → TARGET PLAN (ready for outreach)", len(cols),
                  "Assign anchors per the strategy shares. Exact-match only on genuine local directories, sparingly.")
    header_row(ws, cols, hrow, [5, 52, 46, 16, 9, 44])
    tcolor = {"Branded": AQUA, "Branded+geo": AQUA, "Naked URL": B.ACCENT, "Generic": "7F7F7F",
              "Partial/topical": B.YELLOW, "Partial/natural": B.YELLOW, "Topical/natural": B.YELLOW,
              "Exact-match": RED}
    for i, (u, anc, typ, pri, note) in enumerate(PLAN, 1):
        ws.append([i, u, anc, typ, pri, note])
        tc = ws.cell(row=ws.max_row, column=4)
        tc.fill = fill(tcolor.get(typ, "7F7F7F")); tc.font = F(9, True, WHITE)
        tc.alignment = Alignment(horizontal="center")
        if pri in ("TOP", "High"):
            ws.cell(row=ws.max_row, column=5).font = F(10, True, RED if pri == "TOP" else NAVY)
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(cols))}{ws.max_row}"

    out = os.path.join(BASE, "MasterLawn_Anchor_Target_Plan.xlsx")
    wb.save(out)
    print("saved", out, "| anchors:", len(PLAN), "| targets:", len(TARGETS))

if __name__ == "__main__":
    build()
