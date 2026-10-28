import sys, openpyxl
from openpyxl.utils import get_column_letter

tmpl = openpyxl.load_workbook(sys.argv[1])
out = openpyxl.load_workbook(sys.argv[2])
issues = []

def cellstyle(c):
    f=c.font; fill=c.fill; b=c.border; al=c.alignment
    fg=fill.fgColor.rgb if fill and fill.patternType else None
    bs=tuple(sorted(s for s in ['left','right','top','bottom'] if getattr(b,s) and getattr(b,s).style))
    return (f.name,f.sz,f.bold,(f.color.rgb if f.color else None),fill.patternType,fg,bs,al.horizontal,al.vertical,al.wrap_text)

# tab names/order
if tmpl.sheetnames != out.sheetnames:
    issues.append(f"TAB ORDER: tmpl={tmpl.sheetnames} out={out.sheetnames}")

for name in tmpl.sheetnames:
    if name not in out.sheetnames:
        continue
    t=tmpl[name]; o=out[name]
    # headers
    if name=='Summary':
        # compare header row 3 and title styling
        for coord in ['B1','A3','A4']:
            if cellstyle(t[coord])!=cellstyle(o[coord]):
                issues.append(f"[{name}] {coord} style: tmpl={cellstyle(t[coord])} out={cellstyle(o[coord])}")
        continue
    # data tabs: header row 1
    ncol=t.max_column
    for c in range(1,ncol+1):
        L=get_column_letter(c)
        tv=t.cell(1,c).value; ov=o.cell(1,c).value
        if tv!=ov:
            issues.append(f"[{name}] header {L}1: tmpl={tv!r} out={ov!r}")
        if cellstyle(t.cell(1,c))!=cellstyle(o.cell(1,c)):
            issues.append(f"[{name}] header {L}1 STYLE: tmpl={cellstyle(t.cell(1,c))} out={cellstyle(o.cell(1,c))}")
        # width
        tw=t.column_dimensions[L].width; ow=o.column_dimensions[L].width
        if tw and ow and abs(tw-ow)>0.3:
            issues.append(f"[{name}] width {L}: tmpl={tw} out={ow}")
    # data cell style (row 2)
    if o.max_row>=2 and t.max_row>=2:
        for c in range(1,min(ncol,4)+1):
            if cellstyle(t.cell(2,c))!=cellstyle(o.cell(2,c)):
                issues.append(f"[{name}] data {get_column_letter(c)}2 STYLE: tmpl={cellstyle(t.cell(2,c))} out={cellstyle(o.cell(2,c))}")
    # freeze / filter / gridlines
    if t.freeze_panes!=o.freeze_panes:
        issues.append(f"[{name}] freeze: tmpl={t.freeze_panes} out={o.freeze_panes}")
    if (t.auto_filter.ref is None)!=(o.auto_filter.ref is None):
        issues.append(f"[{name}] autofilter presence: tmpl={t.auto_filter.ref} out={o.auto_filter.ref}")
    if bool(t.sheet_view.showGridLines)!=bool(o.sheet_view.showGridLines):
        issues.append(f"[{name}] gridlines: tmpl={t.sheet_view.showGridLines} out={o.sheet_view.showGridLines}")
    # header height
    th=t.row_dimensions[1].height; oh=o.row_dimensions[1].height
    if th and oh and abs(th-oh)>0.3:
        issues.append(f"[{name}] header height: tmpl={th} out={oh}")
    # data row height
    if t.max_row>=2 and o.max_row>=2:
        th2=t.row_dimensions[2].height; oh2=o.row_dimensions[2].height
        if th2 and oh2 and abs(th2-oh2)>0.3:
            issues.append(f"[{name}] data row height: tmpl={th2} out={oh2}")

# logo presence on Summary
ti=len(getattr(tmpl['Summary'],'_images',[])); oi=len(getattr(out['Summary'],'_images',[]))
if ti!=oi:
    issues.append(f"Summary logo count: tmpl={ti} out={oi}")

if not issues:
    print("STRUCTURAL MATCH: zero differences")
else:
    print(f"{len(issues)} structural differences:")
    for i in issues:
        print("  -", i)
