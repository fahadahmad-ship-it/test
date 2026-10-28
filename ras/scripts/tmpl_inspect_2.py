import sys, openpyxl
from openpyxl.utils import get_column_letter

path = sys.argv[1]
wb = openpyxl.load_workbook(path)
ws = wb['Summary']
print("=== SUMMARY full content ===")
for r in range(1, ws.max_row+1):
    for c in range(1, ws.max_column+1):
        v = ws.cell(r,c).value
        if v is not None:
            print(f"  {get_column_letter(c)}{r} = {v!r}")

def style(cell):
    f = cell.font
    fill = cell.fill
    b = cell.border
    al = cell.alignment
    fg = fill.fgColor.rgb if fill and fill.patternType else None
    bsides = {s: getattr(b,s).style for s in ['left','right','top','bottom'] if getattr(b,s) and getattr(b,s).style}
    return (f"font={f.name} sz={f.sz} bold={f.bold} color={f.color.rgb if f.color else None} | "
            f"fill={fill.patternType}:{fg} | border={bsides} | "
            f"align(h={al.horizontal},v={al.vertical},wrap={al.wrap_text})")

print("\n=== STYLING SAMPLES ===")
print("Summary B1 (title):", style(ws['B1']))
print("Summary A3 (status hdr):", style(ws['A3']))
print("Summary A4 (status data):", style(ws['A4']))
for name in ['Referring Domain Audit','All Backlinks','Exact Match Anchors']:
    w = wb[name]
    print(f"\n[{name}] A1 header:", style(w['A1']))
    print(f"[{name}] A2 data:", style(w['A2']))
    # find an input cell on Exact Match (F/G)
if 'Exact Match Anchors' in wb.sheetnames:
    w=wb['Exact Match Anchors']
    print("ExactMatch F2 (input?):", style(w['F2']))
    print("ExactMatch G2 (input?):", style(w['G2']))

print("\n=== IMAGES ===")
for ws in wb.worksheets:
    imgs = getattr(ws, '_images', [])
    for im in imgs:
        a = im.anchor
        try:
            frm = a._from
            print(f"  {ws.title}: image w={im.width} h={im.height} fromCol={frm.col} colOff={frm.colOff} fromRow={frm.row} rowOff={frm.rowOff}")
            ext = getattr(a, 'ext', None)
            if ext: print(f"     ext cx={ext.cx} cy={ext.cy}")
        except Exception as e:
            print("   img anchor", a, e)
